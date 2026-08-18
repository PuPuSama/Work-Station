from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
import threading
import time
import unittest
import uuid
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import sqlalchemy as sa
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from PIL import Image


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from knowledge_agent.database import create_knowledge_engine  # noqa: E402
from knowledge_agent.assets import (  # noqa: E402
    KnowledgeAsset,
    PostgresKnowledgeAssetRepository,
)
from knowledge_agent.object_storage import (  # noqa: E402
    ProjectKnowledgeObjectService,
)
from knowledge_agent.schema import (  # noqa: E402
    knowledge_assets,
    knowledge_chunks,
    knowledge_product_asset_evidence,
    knowledge_product_source_evidence,
    knowledge_products,
    knowledge_sources,
    projects,
    snapshot_assets,
    source_snapshots,
)
from models import (  # noqa: E402
    PromptSnapshot,
    SeoReviewChange,
    SeoReviewDimension,
    SeoReviewRun,
    TaskRecord,
)
from server_schema import (  # noqa: E402
    article_tasks,
    background_jobs,
    job_batches,
    organizations,
    project_memberships,
    project_ownership,
    task_intakes,
    task_store_state,
    workspace_users,
)
from services.postgres_task_repository import (  # noqa: E402
    PostgresTaskRepository,
)
from services.server_auth import (  # noqa: E402
    SERVER_AUTH_COOKIE_NAME,
    ServerActorSessionCodec,
)
from services.access_control import (  # noqa: E402
    ActorIdentity,
    PostgresProjectAccessRepository,
    ProjectAccessService,
)
from services.object_store import (  # noqa: E402
    StoredObject,
    build_project_object_key,
)
from services.job_queue import JobCancelled, JobConflict  # noqa: E402
from services.postgres_job_queue import PostgresJobQueue  # noqa: E402
from services.server_product_rediscovery import (  # noqa: E402
    ProductRediscoveryCommand,
    ProductRediscoveryUnavailable,
    ServerProductRediscoveryHandler,
    ServerProductRediscoveryRegistry,
)
from services.server_product_selection import (  # noqa: E402
    PostgresConfirmedProductSelection,
)
from services.server_outline_generation import (  # noqa: E402
    PostgresPublishedOutlineContext,
    ProjectPromptReference,
    ServerOutlineGenerationHandler,
    ServerOutlineGenerationRegistry,
)
from services.server_title_generation import (  # noqa: E402
    ServerTitleGenerationHandler,
    ServerTitleGenerationRegistry,
    TitleTemplateReference,
)
from services.server_article_generation import (  # noqa: E402
    ServerArticleGenerationHandler,
    ServerArticleGenerationRegistry,
)
from services.server_link_restoration import (  # noqa: E402
    LinkTemplateReference,
    ServerLinkRestorationHandler,
    ServerLinkRestorationRegistry,
)
from services.server_seo_review_generation import (  # noqa: E402
    ReviewTemplateReference,
    ServerSeoReviewGenerationHandler,
    ServerSeoReviewGenerationRegistry,
)
from services.server_humanize_generation import (  # noqa: E402
    HumanizeGenerationUnavailable,
    ServerHumanizeGenerationHandler,
    ServerHumanizeGenerationRegistry,
)
from services.zerogpt import ZeroGPTDetectionResult  # noqa: E402
from services.server_project_prompts import (  # noqa: E402
    PostgresProjectPromptService,
)
from services.server_project_tasks import (  # noqa: E402
    ServerProjectTaskStoreFactory,
)
from services.server_task_commands import (  # noqa: E402
    ServerTaskCommandUnavailable,
)
from services.server_task_intake import (  # noqa: E402
    PostgresServerTaskIntakeService,
    ServerTaskIntakeResult,
    ServerTaskIntakeRow,
)
from services.server_task_writing_settings import (  # noqa: E402
    ServerTaskWritingSettingsServiceFactory,
)
from storage import content_hash  # noqa: E402
from services.seo_review import (  # noqa: E402
    effective_review_prompt_snapshot,
    parse_seo_review_response,
)


SERVER_ARTICLE = """# Example Buyer Guide

This introduction points readers to [example.com](https://example.com/) before the detailed guidance.

## Buyer Checks

### Confirm the application

Keep the original application guidance.

### Compare evidence

Keep the original evidence guidance.

## FAQ

### What should buyers send?

Send requirements and quantities.

### When should buyers request samples?

Request samples before approval.

### Why compare supplier capability?

Capability affects quality and support.
"""

SERVER_TDK_RESPONSE = """{
  "description": "Compare buyer requirements, product evidence, and supplier capability for a more reliable B2B sourcing decision.",
  "keywords": [
    "buyer requirements",
    "product evidence",
    "supplier capability",
    "B2B sourcing",
    "quality checks",
    "purchase planning"
  ]
}"""


class StubServerTdkLlm:
    ready = True

    def __init__(self, *args, **kwargs) -> None:
        del args, kwargs

    def chat(self, messages, temperature=0.7, max_tokens=1800):
        del messages, temperature, max_tokens
        return SERVER_TDK_RESPONSE


class RecordingOutlineProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def generate(
        self,
        task,
        *,
        prompt_snapshot,
        context_chunks,
    ):
        self.calls.append(
            {
                "task_id": task.id,
                "prompt_id": prompt_snapshot.prompt_id,
                "prompt_version": prompt_snapshot.version,
                "prompt_source": prompt_snapshot.source,
                "chunk_ids": [
                    chunk.chunk_id for chunk in context_chunks
                ],
                "chunk_text": [
                    chunk.text for chunk in context_chunks
                ],
            }
        )
        return "## Generated server outline\n\n### Evidence-led section"


class RecordingTitleProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def generate(
        self,
        task,
        *,
        title_count,
        context_chunks,
    ):
        self.calls.append(
            {
                "task_id": task.id,
                "title_count": title_count,
                "chunk_ids": [
                    chunk.chunk_id for chunk in context_chunks
                ],
            }
        )
        return tuple(
            f"Server candidate title {index}"
            for index in range(1, title_count + 1)
        )


class RecordingArticleProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def generate(
        self,
        task,
        *,
        target_words,
        prompt_snapshot,
        context_chunks,
    ):
        self.calls.append(
            {
                "task_id": task.id,
                "target_words": target_words,
                "prompt_source": prompt_snapshot.source,
                "prompt_version": prompt_snapshot.version,
                "chunk_ids": [
                    chunk.chunk_id for chunk in context_chunks
                ],
                "chunk_text": [
                    chunk.text for chunk in context_chunks
                ],
            }
        )
        return SERVER_ARTICLE


class RecordingAiRateDetector:
    ready = True

    def __init__(self) -> None:
        self.calls: list[str] = []

    def detect(self, text: str) -> ZeroGPTDetectionResult:
        self.calls.append(text)
        return ZeroGPTDetectionResult(
            ai_percentage=18.5,
            ai_words=37,
            text_words=200,
        )


class RecordingLinkRestorationProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def restore(
        self,
        *,
        source_article,
        candidate_article,
        missing_links,
    ):
        self.calls.append(
            {
                "source_hash": content_hash(source_article),
                "candidate_hash": content_hash(candidate_article),
                "missing_links": list(missing_links),
            }
        )
        return source_article


class RecordingSeoReviewProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def generate(
        self,
        task,
        *,
        article,
        prompt_snapshot,
        context_chunks,
    ):
        self.calls.append(
            {
                "task_id": task.id,
                "article_hash": content_hash(article),
                "prompt_source": prompt_snapshot.source,
                "chunk_ids": [
                    chunk.chunk_id for chunk in context_chunks
                ],
                "chunk_text": [
                    chunk.text for chunk in context_chunks
                ],
            }
        )
        return parse_seo_review_response(
            json.dumps(
                {
                    "publish_ready": True,
                    "publish_recommendation": "Ready for human review.",
                    "dimensions": [
                        {
                            "key": "eeat",
                            "name": "E-E-A-T",
                            "score": 9,
                            "target_score": 9,
                            "main_issue": "",
                            "needs_revision": False,
                        }
                    ],
                    "report": "## Review\n\nNo blocking issue.",
                    "changes": [],
                }
            ),
            source_article=article,
            prompt_snapshot=effective_review_prompt_snapshot(
                prompt_snapshot
            ),
            brand_name=task.brand_name,
        )


class RecordingHumanizeProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def generate(self, task, *, source_article, prompt_snapshot):
        self.calls.append(
            {
                "task_id": task.id,
                "source_hash": content_hash(source_article),
                "prompt_id": prompt_snapshot.prompt_id,
                "prompt_version": prompt_snapshot.version,
            }
        )
        return source_article


class FakeDownloadStore:
    def __init__(self) -> None:
        self.signed: list[tuple[str, int]] = []
        self.objects: dict[str, bytes] = {}
        self.put_calls: list[str] = []

    def check_ready(self):
        return None

    def put(self, **kwargs):
        key = str(kwargs["key"])
        data = bytes(kwargs["data"])
        content_type = str(kwargs["content_type"])
        self.objects[key] = data
        self.put_calls.append(key)
        return StoredObject(
            key=key,
            content_hash=hashlib.sha256(data).hexdigest(),
            content_type=content_type,
            byte_size=len(data),
            etag="etag",
        )

    def get(self, key, *, max_bytes):
        return self.objects[key][: max_bytes + 1]

    def create_download_url(self, key, *, expires_seconds):
        self.signed.append((key, expires_seconds))
        return f"https://signed.example.test/{key}"

    def delete(self, key):
        raise AssertionError("not used")


class RecordingAuditWriter:
    def __init__(self) -> None:
        self.events: list[object] = []

    def append(self, connection, event) -> None:
        if not connection.in_transaction():
            raise AssertionError("audit must share the Task transaction")
        self.events.append(event)


class FailingAuditWriter:
    def append(self, connection, event) -> None:
        del connection, event
        raise RuntimeError("private audit failure")


@unittest.skipUnless(
    os.environ.get("ARTICLE_AGENT_DATABASE_URL"),
    "ARTICLE_AGENT_DATABASE_URL is required for PostgreSQL integration tests",
)
class ServerProjectTaskApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.engine = create_knowledge_engine(
            os.environ["ARTICLE_AGENT_DATABASE_URL"]
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.engine.dispose()

    def setUp(self) -> None:
        prefix = f"m7-server-tasks-{uuid.uuid4().hex}"
        self.org_a = f"{prefix}-org-a"
        self.org_b = f"{prefix}-org-b"
        self.user_a = f"{prefix}-user-a"
        self.user_b = f"{prefix}-user-b"
        self.project_a = f"{prefix}.a.example.test"
        self.project_b = f"{prefix}.b.example.test"
        self.task_a = f"{prefix}-task-a"
        self.task_b = f"{prefix}-task-b"
        self.asset_a = f"{prefix}-asset-a"
        self.bad_asset = f"{prefix}-bad-asset"
        self.product_a = f"{prefix}-product-a"
        self.product_b = f"{prefix}-product-b"
        self.inbox_product = f"{prefix}-inbox-product"
        self.unpublished_product = f"{prefix}-unpublished-product"
        with self.engine.begin() as connection:
            connection.execute(
                organizations.insert(),
                (
                    {"organization_id": self.org_a, "name": "Org A"},
                    {"organization_id": self.org_b, "name": "Org B"},
                ),
            )
            connection.execute(
                workspace_users.insert(),
                (
                    {
                        "organization_id": self.org_a,
                        "user_id": self.user_a,
                        "display_name": "User A",
                    },
                    {
                        "organization_id": self.org_b,
                        "user_id": self.user_b,
                        "display_name": "User B",
                    },
                ),
            )
            connection.execute(
                projects.insert(),
                (
                    {
                        "project_id": self.project_a,
                        "customer_name": "Project A",
                        "official_domain": self.project_a,
                    },
                    {
                        "project_id": self.project_b,
                        "customer_name": "Project B",
                        "official_domain": self.project_b,
                    },
                ),
            )
            connection.execute(
                project_ownership.insert(),
                (
                    {
                        "project_id": self.project_a,
                        "organization_id": self.org_a,
                    },
                    {
                        "project_id": self.project_b,
                        "organization_id": self.org_b,
                    },
                ),
            )
            connection.execute(
                project_memberships.insert(),
                (
                    {
                        "organization_id": self.org_a,
                        "project_id": self.project_a,
                        "user_id": self.user_a,
                        "role": "viewer",
                        "granted_by_user_id": self.user_a,
                    },
                    {
                        "organization_id": self.org_b,
                        "project_id": self.project_b,
                        "user_id": self.user_b,
                        "role": "viewer",
                        "granted_by_user_id": self.user_b,
                    },
                ),
            )
        self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        ).upsert(self._task(self.task_a, self.project_a, 2))
        self._task_repository(
            organization_id=self.org_b,
            project_id=self.project_b,
        ).upsert(self._task(self.task_b, self.project_b, 1))
        asset_repository = PostgresKnowledgeAssetRepository(self.engine)
        content_hash = hashlib.sha256(b"asset-a").hexdigest()
        object_key = build_project_object_key(
            self.org_a,
            self.project_a,
            content_hash,
        )
        asset_repository.put_asset(
            KnowledgeAsset(
                project_id=self.project_a,
                asset_id=self.asset_a,
                content_hash=content_hash,
                artifact_uri=f"s3://private-bucket/{object_key}",
                content_type="image/webp",
                byte_size=7,
            )
        )
        bad_hash = hashlib.sha256(b"bad-asset").hexdigest()
        asset_repository.put_asset(
            KnowledgeAsset(
                project_id=self.project_a,
                asset_id=self.bad_asset,
                content_hash=bad_hash,
                artifact_uri=(
                    "s3://private-bucket/organizations/"
                    f"{self.org_b}/projects/{self.project_a}/"
                    f"blobs/{bad_hash[:2]}/{bad_hash}"
                ),
                content_type="image/webp",
                byte_size=9,
            )
        )

    def tearDown(self) -> None:
        with self.engine.begin() as connection:
            connection.execute(
                knowledge_chunks.delete().where(
                    knowledge_chunks.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                knowledge_product_asset_evidence.delete().where(
                    knowledge_product_asset_evidence.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                knowledge_product_source_evidence.delete().where(
                    knowledge_product_source_evidence.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                knowledge_products.delete().where(
                    knowledge_products.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                snapshot_assets.delete().where(
                    snapshot_assets.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                knowledge_assets.delete().where(
                    knowledge_assets.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                source_snapshots.delete().where(
                    source_snapshots.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                knowledge_sources.delete().where(
                    knowledge_sources.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                background_jobs.delete().where(
                    background_jobs.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                job_batches.delete().where(
                    job_batches.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                task_intakes.delete().where(
                    task_intakes.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                article_tasks.delete().where(
                    article_tasks.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                task_store_state.delete().where(
                    task_store_state.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                project_memberships.delete().where(
                    project_memberships.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                project_ownership.delete().where(
                    project_ownership.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                workspace_users.delete().where(
                    workspace_users.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )
            connection.execute(
                projects.delete().where(
                    projects.c.project_id.in_(
                        (self.project_a, self.project_b)
                    )
                )
            )
            connection.execute(
                organizations.delete().where(
                    organizations.c.organization_id.in_(
                        (self.org_a, self.org_b)
                    )
                )
            )

    def _task_repository(
        self,
        *,
        organization_id: str,
        project_id: str,
    ) -> PostgresTaskRepository:
        return PostgresTaskRepository(
            self.engine,
            organization_id=organization_id,
            project_id=project_id,
        )

    def _install_recording_audit(
        self,
        application,
        config,
    ) -> RecordingAuditWriter:
        audit = RecordingAuditWriter()
        application.state.server_project_task_store_factory = (
            ServerProjectTaskStoreFactory(
                self.engine,
                config,
                audit=audit,
            )
        )
        return audit

    @staticmethod
    def _task(task_id: str, project_id: str, topic_index: int) -> dict:
        return TaskRecord(
            id=task_id,
            week_folder="server",
            customer=project_id,
            topic_index=topic_index,
            topic=f"Topic {topic_index}",
            status="title_selected",
            selected_title=f"Selected topic {topic_index}",
            task_dir=f"/server/{task_id}",
            created_at="2026-07-30T00:00:00+00:00",
            updated_at="2026-07-30T00:00:00+00:00",
        ).model_dump(mode="json")

    @staticmethod
    def _writing_settings_payload(
        *,
        revision: int = 0,
        kind: str | None = None,
    ) -> dict[str, object]:
        payload: dict[str, object] = {
            "revision": revision,
            "topic_notes": "Private topic notes",
            "outline_custom_prompt": "Private outline instructions",
            "article_custom_prompt": "Private article instructions",
            "use_outline_custom_prompt": True,
            "use_article_custom_prompt": False,
            "outline_prompt_selection": "project_default",
            "article_prompt_selection": "project_default",
            "include_project_introduction": True,
            "include_project_notes": False,
            "include_topic_notes": True,
        }
        if kind is not None:
            payload["kind"] = kind
        return payload

    @staticmethod
    def _image_bytes(color: str) -> bytes:
        output = BytesIO()
        Image.new("RGB", (320, 240), color).save(
            output,
            format="PNG",
        )
        return output.getvalue()

    def _prepare_link_task(
        self,
    ) -> tuple[PostgresTaskRepository, str, str]:
        source = SERVER_ARTICLE.strip()
        candidate = source.replace(
            "[example.com](https://example.com/)",
            "example.com",
        )
        source_hash = content_hash(source)
        candidate_hash = content_hash(candidate)
        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "final_ai_checked",
                "initial_article": source,
                "initial_article_hash": source_hash,
                "humanized_article": candidate,
                "humanized_article_hash": candidate_hash,
                "article": candidate,
                "final_ai_check": {
                    "confirmed": True,
                    "article_hash": candidate_hash,
                },
                "source_links": [
                    {
                        "anchor": "example.com",
                        "url": "https://example.com/",
                        "count": 1,
                    }
                ],
                "article_versions": [],
            }
        )
        repository.upsert(record)
        return repository, source, candidate

    def _prepare_humanize_task(
        self,
    ) -> tuple[PostgresTaskRepository, str]:
        article = SERVER_ARTICLE.strip()
        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "initial_ai_checked",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "humanized_article": "",
                "humanized_article_hash": "",
                "article_versions": [],
            }
        )
        repository.upsert(record)
        return repository, article

    def _store_outline_context(
        self,
        *,
        project_id: str,
        suffix: str,
        text: str,
        status: str = "published",
        source_kind: str = "knowledge_page",
        trust_tier: str = "hard_fact",
    ) -> str:
        source_id = f"{suffix}-source"
        snapshot_id = f"{suffix}-snapshot"
        chunk_id = f"{snapshot_id}:chunk-0"
        with self.engine.begin() as connection:
            connection.execute(
                knowledge_sources.insert().values(
                    project_id=project_id,
                    source_id=source_id,
                    display_name=f"{suffix} source",
                    source_kind=source_kind,
                    trust_tier=trust_tier,
                    status=status,
                    public_source=True,
                    canonical_url=f"https://{project_id}/{suffix}",
                    current_snapshot_id=(
                        snapshot_id if status == "published" else None
                    ),
                )
            )
            connection.execute(
                source_snapshots.insert().values(
                    project_id=project_id,
                    snapshot_id=snapshot_id,
                    source_id=source_id,
                    content_hash=hashlib.sha256(
                        text.encode("utf-8")
                    ).hexdigest(),
                    parser_name="m7-outline-test",
                    parser_version="1",
                    fetched_at=datetime.now(timezone.utc),
                )
            )
            connection.execute(
                knowledge_chunks.insert().values(
                    project_id=project_id,
                    chunk_id=chunk_id,
                    source_id=source_id,
                    snapshot_id=snapshot_id,
                    ordinal=0,
                    heading_path=["Published facts"],
                    text=text,
                    locator={"section": "facts"},
                    metadata={},
                )
            )
        return chunk_id

    def _store_selectable_product(
        self,
        *,
        project_id: str,
        product_id: str,
        asset_id: str | None = None,
        status: str = "confirmed",
        source_status: str = "published",
        manual_specification_tables: list[dict[str, object]] | None = None,
    ) -> None:
        source_id = f"{product_id}-source"
        snapshot_id = f"{product_id}-snapshot"
        canonical_url = f"https://{project_id}/products/{product_id}"
        with self.engine.begin() as connection:
            connection.execute(
                knowledge_sources.insert().values(
                    project_id=project_id,
                    source_id=source_id,
                    display_name=f"{product_id} source",
                    source_kind="product_detail",
                    trust_tier="hard_fact",
                    status=source_status,
                    public_source=True,
                    canonical_url=canonical_url,
                    current_snapshot_id=snapshot_id,
                )
            )
            connection.execute(
                source_snapshots.insert().values(
                    project_id=project_id,
                    snapshot_id=snapshot_id,
                    source_id=source_id,
                    content_hash=hashlib.sha256(
                        product_id.encode("utf-8")
                    ).hexdigest(),
                    parser_name="m7-test",
                    parser_version="1",
                    fetched_at=datetime.now(timezone.utc),
                )
            )
            connection.execute(
                knowledge_products.insert().values(
                    project_id=project_id,
                    product_id=product_id,
                    name=f"Product {product_id}",
                    status=status,
                    canonical_url=canonical_url,
                    category_path=["Fasteners", "Anchors"],
                    metadata={
                        "description": "Published product description.",
                        "main_content_facts": [
                            "Fact one.",
                            "Fact two.",
                        ],
                        "specification_tables": [
                            {
                                "headers": ["Property", "Value"],
                                "rows": [
                                    ["Material", "316 stainless steel"],
                                    ["Length", "50 mm"],
                                ],
                            }
                        ],
                        **(
                            {
                                "manual_specification_tables": (
                                    manual_specification_tables
                                )
                            }
                            if manual_specification_tables is not None
                            else {}
                        ),
                    },
                )
            )
            connection.execute(
                knowledge_product_source_evidence.insert().values(
                    project_id=project_id,
                    product_id=product_id,
                    source_id=source_id,
                    snapshot_id=snapshot_id,
                    relation="primary_detail",
                    confidence=0.99,
                    reason="M7 server selection test evidence",
                    metadata={
                        "selection_projection": {
                            "schema_version": 1,
                            "name": f"Product {product_id}",
                            "canonical_url": canonical_url,
                            "description": "Published product description.",
                            "reference_facts": [
                                "Fact one.",
                                "Fact two.",
                            ],
                            "specification_tables": [
                                {
                                    "headers": ["Property", "Value"],
                                    "rows": [
                                        [
                                            "Material",
                                            "316 stainless steel",
                                        ],
                                        ["Length", "50 mm"],
                                    ],
                                }
                            ],
                        }
                    },
                )
            )
            if asset_id is not None:
                connection.execute(
                    snapshot_assets.insert().values(
                        project_id=project_id,
                        source_id=source_id,
                        snapshot_id=snapshot_id,
                        asset_id=asset_id,
                        evidence_kind="gallery",
                        ordinal=0,
                        source_url=f"{canonical_url}/image.webp",
                    )
                )
                connection.execute(
                    knowledge_product_asset_evidence.insert().values(
                        project_id=project_id,
                        product_id=product_id,
                        source_id=source_id,
                        snapshot_id=snapshot_id,
                        asset_id=asset_id,
                        role="primary",
                        confidence=0.95,
                        reason="Official product gallery primary image",
                    )
                )

    def test_server_task_reads_are_authorized_and_project_scoped(
        self,
    ) -> None:
        import app as app_module

        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                self.assertEqual(
                    client.get("/api/projects").status_code,
                    401,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/tasks"
                    ).status_code,
                    401,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                download_store = FakeDownloadStore()
                app_module.app.state.server_project_object_service = (
                    ProjectKnowledgeObjectService(
                        store=download_store,  # type: ignore[arg-type]
                        bucket="private-bucket",
                        repository=PostgresKnowledgeAssetRepository(
                            self.engine
                        ),
                        access=ProjectAccessService(
                            PostgresProjectAccessRepository(self.engine)
                        ),
                    )
                )
                directory = client.get("/api/projects")
                self.assertEqual(directory.status_code, 200)
                self.assertEqual(
                    directory.json(),
                    [
                        {
                            "project_id": self.project_a,
                            "customer_name": "Project A",
                            "official_domain": self.project_a,
                            "revision": 0,
                            "effective_role": "viewer",
                            "owning_team_id": None,
                            "owner_user_id": None,
                            "is_project_owner": False,
                            "assignment_status": "pending",
                        }
                    ],
                )
                response = client.get(
                    f"/api/projects/{self.project_a}/tasks"
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    [item["id"] for item in response.json()],
                    [self.task_a],
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/tasks/{self.task_a}"
                    ).json()["id"],
                    self.task_a,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/tasks/{self.task_b}"
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_b}/tasks"
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.get("/api/tasks").status_code,
                    404,
                )
                download = client.get(
                    f"/api/projects/{self.project_a}/assets/"
                    f"{self.asset_a}/download",
                    params={"expires_seconds": 120},
                )
                self.assertEqual(download.status_code, 200)
                self.assertEqual(
                    download.json()["asset_id"],
                    self.asset_a,
                )
                self.assertEqual(
                    download_store.signed,
                    [
                        (
                            build_project_object_key(
                                self.org_a,
                                self.project_a,
                                hashlib.sha256(b"asset-a").hexdigest(),
                            ),
                            120,
                        )
                    ],
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/assets/"
                        "missing/download"
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/assets/"
                        f"{self.bad_asset}/download"
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_b}/assets/"
                        f"{self.asset_a}/download"
                    ).status_code,
                    403,
                )
                rewrite_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/rewrite-from-scratch"
                )
                self.assertEqual(
                    client.post(
                        rewrite_path,
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                rewritten = client.post(
                    rewrite_path,
                    json={"revision": 0},
                )
                self.assertEqual(rewritten.status_code, 200)
                self.assertEqual(rewritten.json()["revision"], 1)
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.task.rewritten"],
                )
                self.assertEqual(
                    client.post(
                        rewrite_path,
                        json={"revision": 0},
                    ).status_code,
                    409,
                )
                self.assertFalse(local_state.exists())

                with self.engine.begin() as connection:
                    connection.execute(
                        projects.update()
                        .where(projects.c.project_id == self.project_a)
                        .values(status="archived")
                    )
                self.assertEqual(
                    client.get("/api/projects").json(),
                    [],
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/tasks"
                    ).status_code,
                    403,
                )

                with self.engine.begin() as connection:
                    connection.execute(
                        projects.update()
                        .where(projects.c.project_id == self.project_a)
                        .values(status="active")
                    )
                    connection.execute(
                        workspace_users.update()
                        .where(
                            workspace_users.c.organization_id
                            == self.org_a,
                            workspace_users.c.user_id == self.user_a,
                        )
                        .values(status="disabled")
                    )
                self.assertEqual(
                    client.get("/api/projects").status_code,
                    401,
                )

    def test_server_task_create_and_import_are_scoped_idempotent_and_audited(
        self,
    ) -> None:
        import app as app_module

        context_repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        context_record = context_repository.get(self.task_a)
        assert context_record is not None
        context_record.update(
            {
                "project_introduction": "Shared project introduction",
            }
        )
        context_repository.upsert(context_record)

        with self.engine.begin() as connection:
            connection.execute(
                projects.update()
                .where(projects.c.project_id == self.project_a)
                .values(project_notes="Shared operator project rules")
            )
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )

        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                create_body = {
                    "intake_id": "manual-request-001",
                    "topic": "How to compare forged fastener suppliers",
                    "primary_keyword": "forged fastener suppliers",
                    "competitor_keyword": "forged fasteners",
                    "competitor_blog": (
                        "https://competitor.example.test/guide"
                    ),
                }
                created = client.post(
                    f"/api/projects/{self.project_a}/tasks",
                    json=create_body,
                )
                self.assertEqual(created.status_code, 200, created.text)
                create_payload = created.json()
                self.assertTrue(create_payload["created"])
                self.assertEqual(create_payload["intake_kind"], "manual")
                self.assertEqual(create_payload["source_name"], "manual")
                self.assertRegex(
                    create_payload["source_digest"],
                    r"^[0-9a-f]{64}$",
                )
                self.assertEqual(len(create_payload["tasks"]), 1)
                manual_task = create_payload["tasks"][0]
                self.assertEqual(manual_task["topic_index"], 3)
                self.assertEqual(manual_task["status"], "new")
                self.assertEqual(manual_task["revision"], 0)
                self.assertEqual(
                    set(manual_task),
                    {
                        "id",
                        "topic_index",
                        "topic",
                        "primary_keyword",
                        "competitor_keyword",
                        "competitor_blog",
                        "status",
                        "revision",
                    },
                )

                repository = self._task_repository(
                    organization_id=self.org_a,
                    project_id=self.project_a,
                )
                stored = repository.get(manual_task["id"])
                assert stored is not None
                self.assertEqual(stored["organization_id"], self.org_a)
                self.assertEqual(stored["project_id"], self.project_a)
                self.assertEqual(stored["customer"], self.project_a)
                self.assertEqual(stored["brand_name"], "Project A")
                self.assertEqual(stored["week_folder"], "server")
                self.assertEqual(stored["task_dir"], "")
                self.assertEqual(stored["source_kind"], "manual")
                self.assertEqual(
                    stored["project_introduction"],
                    "Shared project introduction",
                )
                self.assertEqual(
                    stored["project_notes"],
                    "Shared operator project rules",
                )
                self.assertEqual(
                    stored["primary_keyword"],
                    "forged fastener suppliers",
                )
                self.assertNotIn(
                    str(local_state),
                    json.dumps(stored),
                )

                retried = client.post(
                    f"/api/projects/{self.project_a}/tasks",
                    json=create_body,
                )
                self.assertEqual(retried.status_code, 200)
                self.assertFalse(retried.json()["created"])
                self.assertEqual(
                    retried.json()["tasks"],
                    create_payload["tasks"],
                )
                self.assertEqual(len(audit.events), 1)

                conflict = client.post(
                    f"/api/projects/{self.project_a}/tasks",
                    json={**create_body, "topic": "Different input"},
                )
                self.assertEqual(conflict.status_code, 409)
                self.assertNotIn("Different input", conflict.text)

                workbook = Workbook()
                worksheet = workbook.active
                worksheet.append(
                    [
                        "文章话题",
                        "目标关键词",
                        "竞对关键词",
                        "竞对 Blog URL",
                    ]
                )
                worksheet.append(
                    [
                        "Bolt grade selection",
                        "bolt grade guide",
                        "bolt grades",
                        "https://competitor.example.test/bolts",
                    ]
                )
                workbook_output = BytesIO()
                workbook.save(workbook_output)
                workbook.close()
                preview = client.post(
                    f"/api/projects/{self.project_a}/task-imports/preview",
                    files={
                        "file": (
                            "topics.xlsx",
                            workbook_output.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    },
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertEqual(
                    preview.json()["mapping"],
                    {
                        "topic": 0,
                        "primary_keyword": 1,
                        "competitor_keyword": 2,
                        "competitor_blog": 3,
                    },
                )

                imported = client.post(
                    f"/api/projects/{self.project_a}/task-imports",
                    json={
                        "intake_id": "import-request-001",
                        "source_name": "q3-topic-rows.tsv",
                        "rows": [
                            {
                                "topic": "Bolt grade selection",
                                "primary_keyword": "bolt grade guide",
                                "competitor_keyword": "bolt grades",
                                "competitor_blog": "",
                            },
                            {
                                "topic": "Fastener coating guide",
                                "primary_keyword": "fastener coatings",
                                "competitor_keyword": "",
                                "competitor_blog": (
                                    "https://competitor.example.test/coatings"
                                ),
                            },
                        ],
                    },
                )
                self.assertEqual(imported.status_code, 200, imported.text)
                import_payload = imported.json()
                self.assertTrue(import_payload["created"])
                self.assertEqual(
                    import_payload["intake_kind"],
                    "row_import",
                )
                self.assertEqual(
                    [item["topic_index"] for item in import_payload["tasks"]],
                    [4, 5],
                )
                source_kinds: list[str] = []
                for item in import_payload["tasks"]:
                    imported_record = repository.get(item["id"])
                    assert imported_record is not None
                    source_kinds.append(str(imported_record["source_kind"]))
                    self.assertEqual(
                        imported_record["project_introduction"],
                        "Shared project introduction",
                    )
                    self.assertEqual(
                        imported_record["project_notes"],
                        "Shared operator project rules",
                    )
                self.assertEqual(
                    source_kinds,
                    ["server_import", "server_import"],
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.task.created",
                        "article.tasks.imported",
                    ],
                )
                audit_json = json.dumps(
                    [event.details for event in audit.events],
                    ensure_ascii=False,
                )
                self.assertNotIn(
                    "forged fastener",
                    audit_json.casefold(),
                )
                self.assertNotIn(
                    "competitor.example.test",
                    audit_json,
                )
                self.assertNotIn(
                    create_payload["source_digest"],
                    audit_json,
                )
                self.assertEqual(
                    audit.events[-1].details,
                    {
                        "intake_kind": "row_import",
                        "task_count": 2,
                        "first_topic_index": 4,
                        "last_topic_index": 5,
                    },
                )
                self.assertEqual(
                    client.post(
                        f"/api/projects/{self.project_b}/task-imports",
                        json={
                            "intake_id": "import-request-002",
                            "source_name": "rows.tsv",
                            "rows": [{"topic": "Cross project"}],
                        },
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.post(
                        f"/api/projects/{self.project_a}/tasks",
                        json={
                            **create_body,
                            "intake_id": "manual-request-002",
                            "task_id": "caller-controlled",
                        },
                    ).status_code,
                    422,
                )
                invalid_url = client.post(
                    f"/api/projects/{self.project_a}/tasks",
                    json={
                        **create_body,
                        "intake_id": "manual-request-003",
                        "competitor_blog": (
                            "https://user:secret@example.test/private"
                        ),
                    },
                )
                self.assertEqual(invalid_url.status_code, 422)
                self.assertNotIn("secret", invalid_url.text)
                self.assertFalse(local_state.exists())

    def test_server_task_intake_rolls_back_when_audit_fails(
        self,
    ) -> None:
        import app as app_module

        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        isolated = replace(
            base_config,
            knowledge_agent_enabled=False,
        )
        with (
            patch.object(app_module, "config", return_value=isolated),
            patch.dict(
                os.environ,
                {
                    "ARTICLE_AGENT_SERVER_MODE": "true",
                    "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                },
                clear=False,
            ),
            TestClient(app_module.app) as client,
        ):
            client.app.state.server_project_task_store_factory = (
                ServerProjectTaskStoreFactory(
                    self.engine,
                    isolated,
                    audit=FailingAuditWriter(),
                )
            )
            client.cookies.set(
                SERVER_AUTH_COOKIE_NAME,
                codec.create(actor),
            )
            response = client.post(
                f"/api/projects/{self.project_a}/tasks",
                json={
                    "intake_id": "audit-failure-001",
                    "topic": "Must roll back",
                },
            )
            self.assertEqual(response.status_code, 503)
            self.assertNotIn("private audit failure", response.text)

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        self.assertEqual(
            [item["id"] for item in repository.load_all()],
            [self.task_a],
        )
        with self.engine.connect() as connection:
            receipt_count = connection.execute(
                sa.select(sa.func.count())
                .select_from(task_intakes)
                .where(
                    task_intakes.c.organization_id == self.org_a,
                    task_intakes.c.project_id == self.project_a,
                )
            ).scalar_one()
        self.assertEqual(receipt_count, 0)

    def test_concurrent_identical_task_intake_creates_one_receipt(
        self,
    ) -> None:
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        audit = RecordingAuditWriter()
        service = PostgresServerTaskIntakeService(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
            audit=audit,
        )
        actor = ActorIdentity(self.org_a, self.user_a)
        barrier = threading.Barrier(2)
        results: list[ServerTaskIntakeResult] = []
        errors: list[BaseException] = []

        def create() -> None:
            try:
                barrier.wait(timeout=10)
                results.append(
                    service.import_rows(
                        actor=actor,
                        intake_id="concurrent-import-001",
                        source_name="concurrent.tsv",
                        rows=(
                            ServerTaskIntakeRow(
                                topic="Concurrent topic",
                            ),
                        ),
                    )
                )
            except BaseException as exc:
                errors.append(exc)

        threads = [threading.Thread(target=create) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)

        self.assertEqual(errors, [])
        self.assertEqual(len(results), 2)
        created = sorted(result.created for result in results)
        self.assertEqual(created, [False, True])
        task_ids = {result.tasks[0].id for result in results}
        self.assertEqual(len(task_ids), 1)
        self.assertEqual(len(audit.events), 1)
        with self.engine.connect() as connection:
            receipt_count = connection.execute(
                sa.select(sa.func.count())
                .select_from(task_intakes)
                .where(
                    task_intakes.c.organization_id == self.org_a,
                    task_intakes.c.project_id == self.project_a,
                )
            ).scalar_one()
        self.assertEqual(receipt_count, 1)

    def test_server_selects_only_current_title_candidate_with_cas(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "revision": 0,
                "status": "titles_ready",
                "title_candidates": [
                    "First candidate",
                    "Second candidate",
                ],
                "selected_title": "",
                "outline": "stale outline",
                "outline_draft": "stale outline draft",
                "article": "stale article",
            }
        )
        repository.upsert(record)
        codec = ServerActorSessionCodec(b"q" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "q" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/selected-title"
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={"revision": 0, "candidate_index": 1},
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "candidate_index": 1,
                            "title": "Caller replacement",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={"revision": 0, "candidate_index": 9},
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_a}/selected-title"
                        ),
                        json={"revision": 0, "candidate_index": 1},
                    ).status_code,
                    403,
                )
                selected = client.put(
                    path,
                    json={"revision": 0, "candidate_index": 1},
                )
                self.assertEqual(selected.status_code, 200, selected.text)
                response = selected.json()
                self.assertEqual(
                    response["selected_title"],
                    "Second candidate",
                )
                self.assertEqual(response["status"], "title_selected")
                self.assertEqual(response["revision"], 1)
                self.assertEqual(response["outline"], "")
                self.assertEqual(response["outline_draft"], "")
                self.assertEqual(response["article"], "")
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.title.selected"],
                )
                self.assertEqual(
                    audit.events[0].details,
                    {
                        "from_revision": 0,
                        "to_revision": 1,
                        "status": "title_selected",
                        "candidate_count": 2,
                        "candidate_index": 1,
                    },
                )
                self.assertNotIn(
                    "Second candidate",
                    str(audit.events[0].details),
                )
                stale = client.put(
                    path,
                    json={"revision": 0, "candidate_index": 0},
                )
                self.assertEqual(stale.status_code, 409)
                self.assertEqual(len(audit.events), 1)
                stored = repository.get(self.task_a)
                assert stored is not None
                self.assertEqual(
                    stored["selected_title"],
                    "Second candidate",
                )
                self.assertEqual(stored["revision"], 1)
                self.assertFalse(local_state.exists())

    def test_server_saves_outline_draft_and_confirmation_with_cas(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "revision": 0,
                "status": "outline_confirmed",
                "selected_title": "Current title",
                "outline": "## Confirmed outline",
                "outline_draft": "## Confirmed outline",
                "article": "stale article",
                "initial_article": "stale initial article",
                "article_versions": [],
            }
        )
        repository.upsert(record)
        codec = ServerActorSessionCodec(b"r" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "r" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline"
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "outline": "## Working draft",
                            "confirmed": False,
                        },
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "outline": "   ",
                            "confirmed": False,
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "outline": "## Working draft",
                            "confirmed": False,
                            "status": "caller-controlled",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_a}/outline"
                        ),
                        json={
                            "revision": 0,
                            "outline": "## Cross-project draft",
                            "confirmed": False,
                        },
                    ).status_code,
                    403,
                )
                draft = client.put(
                    path,
                    json={
                        "revision": 0,
                        "outline": "## Working draft",
                        "confirmed": False,
                    },
                )
                self.assertEqual(draft.status_code, 200, draft.text)
                draft_body = draft.json()
                self.assertEqual(draft_body["revision"], 1)
                self.assertEqual(
                    draft_body["outline"],
                    "## Confirmed outline",
                )
                self.assertEqual(
                    draft_body["outline_draft"],
                    "## Working draft",
                )
                self.assertEqual(
                    draft_body["article"],
                    "stale article",
                )
                self.assertEqual(
                    draft_body["article_versions"][-1]["kind"],
                    "outline_draft",
                )
                confirmed = client.put(
                    path,
                    json={
                        "revision": 1,
                        "outline": "## Final reviewed outline",
                        "confirmed": True,
                    },
                )
                self.assertEqual(
                    confirmed.status_code,
                    200,
                    confirmed.text,
                )
                confirmed_body = confirmed.json()
                self.assertEqual(confirmed_body["revision"], 2)
                self.assertEqual(
                    confirmed_body["outline"],
                    "## Final reviewed outline",
                )
                self.assertEqual(
                    confirmed_body["outline_draft"],
                    "## Final reviewed outline",
                )
                self.assertEqual(
                    confirmed_body["status"],
                    "outline_confirmed",
                )
                self.assertEqual(confirmed_body["article"], "")
                self.assertEqual(
                    confirmed_body["initial_article"],
                    "",
                )
                self.assertEqual(
                    confirmed_body["article_versions"][-1]["kind"],
                    "outline",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.outline.updated",
                        "article.outline.updated",
                    ],
                )
                self.assertEqual(
                    audit.events[0].details,
                    {
                        "from_revision": 0,
                        "to_revision": 1,
                        "status": "outline_confirmed",
                        "confirmed": False,
                        "outline_characters": 16,
                    },
                )
                self.assertEqual(
                    audit.events[1].details,
                    {
                        "from_revision": 1,
                        "to_revision": 2,
                        "status": "outline_confirmed",
                        "confirmed": True,
                        "outline_characters": 25,
                    },
                )
                self.assertNotIn(
                    "Working draft",
                    str(audit.events),
                )
                self.assertNotIn(
                    "Final reviewed outline",
                    str(audit.events),
                )
                stale = client.put(
                    path,
                    json={
                        "revision": 1,
                        "outline": "## Stale outline",
                        "confirmed": True,
                    },
                )
                self.assertEqual(stale.status_code, 409)
                self.assertEqual(len(audit.events), 2)
                stored = repository.get(self.task_a)
                assert stored is not None
                self.assertEqual(stored["revision"], 2)
                self.assertEqual(
                    stored["outline"],
                    "## Final reviewed outline",
                )
                self.assertFalse(local_state.exists())

    def test_server_restores_only_owned_outline_version_to_draft(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "revision": 0,
                "status": "outline_confirmed",
                "selected_title": "Current title",
                "outline": "## Current confirmed outline",
                "outline_draft": "## Current confirmed outline",
                "article": "current downstream article",
                "article_versions": [
                    {
                        "kind": "outline",
                        "content": "## Earlier outline",
                        "source_kind": "manual_confirmed",
                    },
                    {
                        "kind": "article",
                        "content": "private article version",
                        "source_kind": "first_version",
                    },
                ],
            }
        )
        repository.upsert(record)
        codec = ServerActorSessionCodec(b"s" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "s" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline/restore-version"
                )
                self.assertEqual(
                    client.post(
                        path,
                        json={"revision": 0, "version_index": 0},
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "version_index": 0,
                            "outline": "caller replacement",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        path,
                        json={"revision": 0, "version_index": 99},
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.post(
                        path,
                        json={"revision": 0, "version_index": 1},
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_a}/outline/restore-version"
                        ),
                        json={"revision": 0, "version_index": 0},
                    ).status_code,
                    403,
                )
                restored = client.post(
                    path,
                    json={"revision": 0, "version_index": 0},
                )
                self.assertEqual(
                    restored.status_code,
                    200,
                    restored.text,
                )
                body = restored.json()
                self.assertEqual(body["revision"], 1)
                self.assertEqual(
                    body["outline"],
                    "## Current confirmed outline",
                )
                self.assertEqual(
                    body["outline_draft"],
                    "## Earlier outline",
                )
                self.assertEqual(
                    body["article"],
                    "current downstream article",
                )
                self.assertEqual(
                    body["status"],
                    "outline_confirmed",
                )
                self.assertEqual(
                    body["article_versions"][-1]["kind"],
                    "outline_draft",
                )
                self.assertEqual(
                    body["article_versions"][-1]["source_kind"],
                    "restored",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.outline_version.restored"],
                )
                self.assertEqual(
                    audit.events[0].details,
                    {
                        "from_revision": 0,
                        "to_revision": 1,
                        "status": "outline_confirmed",
                        "restored_from": "outline",
                        "version_index": 0,
                    },
                )
                self.assertNotIn(
                    "Earlier outline",
                    str(audit.events),
                )
                stale = client.post(
                    path,
                    json={"revision": 0, "version_index": 0},
                )
                self.assertEqual(stale.status_code, 409)
                self.assertEqual(len(audit.events), 1)
                stored = repository.get(self.task_a)
                assert stored is not None
                self.assertEqual(stored["revision"], 1)
                self.assertEqual(
                    stored["outline_draft"],
                    "## Earlier outline",
                )
                self.assertFalse(local_state.exists())

    @unittest.skip("Local mode was removed from the Server-only application")
    def test_server_task_api_is_not_added_to_local_mode(self) -> None:
        import app as app_module

        previous_mode = getattr(
            app_module.app.state,
            "server_mode_enabled",
            None,
        )
        app_module.app.state.server_mode_enabled = False
        try:
            response = TestClient(app_module.app).get(
                f"/api/projects/{self.project_a}/tasks"
            )
            self.assertEqual(response.status_code, 404)
            self.assertEqual(
                TestClient(app_module.app).get(
                    "/api/projects"
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/assets/"
                    f"{self.asset_a}/download"
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks",
                    json={
                        "intake_id": "local-create-001",
                        "topic": "Local must stay isolated",
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/task-imports",
                    json={
                        "intake_id": "local-import-001",
                        "source_name": "rows.tsv",
                        "rows": [{"topic": "Local must stay isolated"}],
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/rewrite-from-scratch",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/writing-settings",
                    json=self._writing_settings_payload(),
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/writing-settings/preview",
                    json=self._writing_settings_payload(kind="outline"),
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/selected-title",
                    json={"revision": 0, "candidate_index": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-reviews",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-reviews/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/humanize",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/humanize/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-reviews/review-a/preview",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-reviews/review-a/"
                    "changes/change-a",
                    json={"revision": 0, "decision": "rejected"},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline/restore-version",
                    json={"revision": 0, "version_index": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline",
                    json={
                        "revision": 0,
                        "outline": "## Reviewed outline",
                        "confirmed": True,
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/titles",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/titles/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article/rewrite",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article/rewrite/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/restore-links",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/restore-links/jobs/job-a",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/initial-ai/screenshot",
                    params={"revision": 0},
                    files={
                        "file": (
                            "initial.png",
                            self._image_bytes("white"),
                            "image/png",
                        )
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/initial-ai",
                    json={"revision": 0, "confirmed": True},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/initial-ai/"
                    "screenshot/download",
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/humanized-article",
                    json={
                        "revision": 0,
                        "article": SERVER_ARTICLE,
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-review-settings",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/products",
                    json={
                        "revision": 0,
                        "product_ids": [self.product_a],
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).put(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article/sections",
                    json={
                        "revision": 0,
                        "heading_path": ["Buyer Checks"],
                        "replacement_body": "Replacement.",
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/product-rediscovery",
                    json={
                        "revision": 0,
                        "category_url": (
                            f"https://{self.project_a}/products"
                        ),
                    },
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).post(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/export-docx",
                    json={"revision": 0},
                ).status_code,
                404,
            )
            self.assertEqual(
                TestClient(app_module.app).get(
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/docx/download",
                ).status_code,
                404,
            )
        finally:
            app_module.app.state.server_mode_enabled = previous_mode

    def test_server_replaces_products_only_from_confirmed_project_catalog(
        self,
    ) -> None:
        import app as app_module

        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.product_a,
            asset_id=self.asset_a,
        )
        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.inbox_product,
            status="inbox",
        )
        self._store_selectable_product(
            project_id=self.project_b,
            product_id=self.product_b,
        )
        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.unpublished_product,
            source_status="needs_review",
        )
        with self.engine.begin() as connection:
            connection.execute(
                knowledge_products.update()
                .where(
                    knowledge_products.c.project_id == self.project_a,
                    knowledge_products.c.product_id == self.product_a,
                )
                .values(
                    name="Unreviewed replacement name",
                    canonical_url="https://unreviewed.invalid/product",
                    metadata={
                        "description": "Unreviewed replacement description.",
                        "main_content_facts": ["Unreviewed fact."],
                    },
                )
            )
        codec = ServerActorSessionCodec(b"p" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "p" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/products"
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.product_a],
                        },
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.product_b],
                        },
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.unpublished_product],
                        },
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.inbox_product],
                        },
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [
                                self.product_a,
                                self.product_a,
                            ],
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.product_a],
                            "products": [
                                {
                                    "product_id": self.product_a,
                                    "canonical_url": "https://attacker.invalid",
                                }
                            ],
                        },
                    ).status_code,
                    422,
                )

                recommendation_reason = (
                    "Matches the article's corrosion-resistant fastening "
                    "application."
                )
                with self.engine.begin() as connection:
                    task_payload = dict(
                        connection.execute(
                            sa.select(article_tasks.c.payload).where(
                                article_tasks.c.organization_id == self.org_a,
                                article_tasks.c.project_id == self.project_a,
                                article_tasks.c.task_id == self.task_a,
                            )
                        ).scalar_one()
                    )
                    task_payload["product_candidate_ids"] = [self.product_a]
                    task_payload["product_candidate_reasons"] = {
                        self.product_a: recommendation_reason
                    }
                    connection.execute(
                        article_tasks.update()
                        .where(
                            article_tasks.c.organization_id == self.org_a,
                            article_tasks.c.project_id == self.project_a,
                            article_tasks.c.task_id == self.task_a,
                        )
                        .values(payload=task_payload)
                    )

                response = client.put(
                    path,
                    json={
                        "revision": 0,
                        "product_ids": [self.product_a],
                    },
                )
                self.assertEqual(response.status_code, 200, response.text)
                saved = response.json()
                self.assertEqual(saved["revision"], 1)
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.products.confirmed"],
                )
                self.assertEqual(
                    [item["product_id"] for item in saved["products"]],
                    [self.product_a],
                )
                product = saved["products"][0]
                self.assertEqual(product["selection_reason"], recommendation_reason)
                self.assertEqual(saved["product_candidate_ids"], [])
                self.assertEqual(saved["product_candidate_reasons"], {})
                self.assertEqual(
                    product["name"],
                    f"Product {self.product_a}",
                )
                self.assertEqual(
                    product["canonical_url"],
                    f"https://{self.project_a}/products/{self.product_a}",
                )
                self.assertEqual(
                    product["description"],
                    "Published product description.",
                )
                self.assertEqual(product["selected_asset_id"], self.asset_a)
                self.assertEqual(product["image_path"], "")
                self.assertEqual(product["asset_count"], 1)
                self.assertEqual(
                    product["reference_facts"],
                    ["Fact one.", "Fact two."],
                )
                self.assertEqual(
                    product["specifications"],
                    {
                        "Material": "316 stainless steel",
                        "Length": "50 mm",
                    },
                )
                self.assertNotIn("artifact_uri", product)
                self.assertNotIn("source_url", product)
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "product_ids": [self.product_a],
                        },
                    ).status_code,
                    409,
                )
                self.assertFalse(local_state.exists())

    def test_manual_product_specifications_are_used_for_task_selection(
        self,
    ) -> None:
        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.product_a,
            manual_specification_tables=[
                {
                    "headers": ["Parameter", "6000W", "8000W"],
                    "rows": [["Surge Power", "12000VA", "16000VA"]],
                }
            ],
        )

        selected = PostgresConfirmedProductSelection(self.engine).select(
            self.project_a,
            [self.product_a],
        )[0]

        self.assertTrue(selected.specifications_overridden)
        self.assertEqual(
            selected.specifications["Surge Power [8000W]"],
            "16000VA",
        )
        self.assertEqual(
            selected.specifications["Surge Power [6000W]"],
            "12000VA",
        )

    def test_server_catalog_is_minimal_current_and_project_scoped(
        self,
    ) -> None:
        import app as app_module

        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.product_a,
            asset_id=self.asset_a,
        )
        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.inbox_product,
            status="inbox",
        )
        self._store_selectable_product(
            project_id=self.project_a,
            product_id=self.unpublished_product,
            source_status="needs_review",
        )
        self._store_selectable_product(
            project_id=self.project_b,
            product_id=self.product_b,
        )
        codec = ServerActorSessionCodec(b"g" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "g" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                response = client.get(
                    f"/api/projects/{self.project_a}/catalog"
                    f"?image_product_ids={self.product_a}"
                )
                self.assertEqual(response.status_code, 200, response.text)
                body = response.json()
                self.assertEqual(
                    body["products"],
                    [
                        {
                            "asset_count": 1,
                            "name": f"Product {self.product_a}",
                            "product_id": self.product_a,
                            "selected_asset_id": self.asset_a,
                        }
                    ],
                )
                self.assertEqual(len(body["image_assets"]), 1)
                image = body["image_assets"][0]
                self.assertEqual(image["asset_id"], self.asset_a)
                self.assertEqual(image["content_type"], "image/webp")
                self.assertEqual(
                    set(image),
                    {
                        "asset_id",
                        "product_id",
                        "byte_size",
                        "content_type",
                        "evidence_kind",
                        "height",
                        "label",
                        "width",
                    },
                )
                serialized = response.text.casefold()
                for private_field in (
                    "artifact_uri",
                    "object_key",
                    "content_hash",
                    "canonical_url",
                    "source_url",
                    "metadata",
                    "private-bucket",
                ):
                    self.assertNotIn(private_field, serialized)
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_b}/catalog"
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.get(
                        f"/api/projects/{self.project_a}/catalog?image_limit=0"
                    ).status_code,
                    422,
                )

    def test_server_prepares_private_asset_ids_without_local_paths(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.replace(
            "Keep the original application guidance.",
            (
                f"Product {self.product_a} supports the application "
                "described in this section."
            ),
        )
        hero_asset_id = f"{self.asset_a}-hero-source"
        product_asset_id = f"{self.asset_a}-product-source"
        record.update(
            {
                "status": "links_verified",
                "selected_title": "Example Buyer Guide",
                "linked_article": article,
                "article": article,
                "products": [
                    {
                        "product_id": self.product_a,
                        "name": f"Product {self.product_a}",
                        "url": (
                            f"https://{self.project_a}/products/"
                            f"{self.product_a}"
                        ),
                        "canonical_url": (
                            f"https://{self.project_a}/products/"
                            f"{self.product_a}"
                        ),
                        "selected_asset_id": product_asset_id,
                        "asset_status": "ready",
                    }
                ],
                "article_versions": [],
            }
        )
        repository.upsert(record)

        codec = ServerActorSessionCodec(b"i" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        private_store = FakeDownloadStore()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        object_service = ProjectKnowledgeObjectService(
            store=private_store,
            bucket="private-bucket",
            repository=PostgresKnowledgeAssetRepository(self.engine),
            access=access,
        )
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "i" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/prepare-images"
                )
                payload = {
                    "revision": 0,
                    "hero_asset_id": hero_asset_id,
                    "product_asset_ids": {
                        self.product_a: product_asset_id,
                    },
                }
                self.assertEqual(
                    client.post(path, json=payload).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                object_service.upload(
                    actor=actor,
                    project_id=self.project_a,
                    asset_id=hero_asset_id,
                    data=self._image_bytes("navy"),
                    content_type="image/png",
                    width=320,
                    height=240,
                )
                object_service.upload(
                    actor=actor,
                    project_id=self.project_a,
                    asset_id=product_asset_id,
                    data=self._image_bytes("orange"),
                    content_type="image/png",
                    width=320,
                    height=240,
                )
                self._store_selectable_product(
                    project_id=self.project_a,
                    product_id=self.product_a,
                    asset_id=product_asset_id,
                )
                with self.engine.begin() as connection:
                    source_id = f"{self.product_a}-source"
                    snapshot_id = f"{self.product_a}-snapshot"
                    connection.execute(
                        snapshot_assets.insert().values(
                            project_id=self.project_a,
                            source_id=source_id,
                            snapshot_id=snapshot_id,
                            asset_id=hero_asset_id,
                            evidence_kind="gallery",
                            ordinal=1,
                            source_url=(
                                f"https://{self.project_a}/products/"
                                f"{self.product_a}/hero.png"
                            ),
                        )
                    )
                    connection.execute(
                        knowledge_product_asset_evidence.insert().values(
                            project_id=self.project_a,
                            product_id=self.product_a,
                            source_id=source_id,
                            snapshot_id=snapshot_id,
                            asset_id=hero_asset_id,
                            role="hero",
                            confidence=0.95,
                            reason="Selected product hero image",
                        )
                    )
                client.app.state.server_project_object_service = (
                    object_service
                )

                rejected_hero = client.post(
                    path,
                    json={
                        **payload,
                        "hero_asset_id": "asset-from-another-product",
                    },
                )
                self.assertEqual(
                    rejected_hero.status_code,
                    422,
                    rejected_hero.text,
                )
                self.assertIn(
                    "must belong to a selected product",
                    rejected_hero.text,
                )

                rejected = client.post(
                    path,
                    json={
                        **payload,
                        "product_asset_ids": {
                            self.product_a: "asset-from-another-product",
                        },
                    },
                )
                self.assertEqual(rejected.status_code, 422, rejected.text)
                self.assertIn(
                    "must belong to their selected products",
                    rejected.text,
                )

                prepared = client.post(path, json=payload)

                self.assertEqual(
                    prepared.status_code,
                    200,
                    prepared.text,
                )
                saved = prepared.json()
                self.assertEqual(saved["revision"], 1)
                self.assertEqual(saved["status"], "images_ready")
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.images.prepared"],
                )
                self.assertEqual(len(saved["images"]), 2)
                self.assertEqual(
                    [
                        item["source_asset_id"]
                        for item in saved["images"]
                    ],
                    [hero_asset_id, product_asset_id],
                )
                for image in saved["images"]:
                    self.assertEqual(image["source_path"], "")
                    self.assertEqual(image["prepared_path"], "")
                    self.assertTrue(image["prepared_asset_id"])
                    self.assertNotIn("artifact_uri", image)
                    self.assertNotIn("source_url", image)
                self.assertEqual(
                    saved["images"][0]["anchor_after"],
                    "before_first_h2",
                )
                self.assertEqual(
                    saved["images"][1]["anchor_heading"],
                    "Confirm the application",
                )
                put_count = len(private_store.put_calls)
                self.assertEqual(
                    client.post(path, json=payload).status_code,
                    409,
                )
                self.assertEqual(
                    len(private_store.put_calls),
                    put_count,
                )

                derived_asset_id = saved["images"][1][
                    "prepared_asset_id"
                ]
                downloaded = client.get(
                    (
                        f"/api/projects/{self.project_a}/assets/"
                        f"{derived_asset_id}/download"
                    )
                )
                self.assertEqual(
                    downloaded.status_code,
                    200,
                    downloaded.text,
                )
                self.assertEqual(
                    downloaded.json()["asset_id"],
                    derived_asset_id,
                )

                export_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/export-docx"
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="viewer")
                    )
                self.assertEqual(
                    client.post(
                        export_path,
                        json={"revision": 1},
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                put_count = len(private_store.put_calls)
                exported = client.post(
                    export_path,
                    json={"revision": 1},
                )
                self.assertEqual(
                    exported.status_code,
                    200,
                    exported.text,
                )
                delivered = exported.json()
                self.assertEqual(delivered["revision"], 2)
                self.assertEqual(
                    delivered["status"],
                    "docx_exported",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.images.prepared",
                        "article.docx.exported",
                    ],
                )
                self.assertEqual(delivered["docx_path"], "")
                self.assertTrue(delivered["docx_asset_id"])
                self.assertEqual(
                    len(delivered["docx_content_hash"]),
                    64,
                )
                self.assertTrue(
                    delivered["docx_filename"].endswith(".docx")
                )
                self.assertEqual(
                    len(private_store.put_calls),
                    put_count + 1,
                )
                docx_asset = (
                    PostgresKnowledgeAssetRepository(
                        self.engine
                    ).get_asset(
                        self.project_a,
                        delivered["docx_asset_id"],
                    )
                )
                assert docx_asset is not None
                docx_key = str(docx_asset.metadata["object_key"])
                with ZipFile(
                    BytesIO(private_store.objects[docx_key])
                ) as archive:
                    self.assertIn(
                        "word/document.xml",
                        archive.namelist(),
                    )
                    self.assertEqual(
                        len(
                            [
                                name
                                for name in archive.namelist()
                                if name.startswith("word/media/")
                            ]
                        ),
                        2,
                    )
                self.assertEqual(
                    client.get(
                        (
                            f"/api/projects/{self.project_a}/assets/"
                            f"{delivered['docx_asset_id']}/download"
                        )
                    ).status_code,
                    404,
                )
                docx_download_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/docx/download"
                )
                docx_download = client.get(docx_download_path)
                self.assertEqual(
                    docx_download.status_code,
                    200,
                    docx_download.text,
                )
                self.assertEqual(
                    docx_download.json()["asset_id"],
                    delivered["docx_asset_id"],
                )
                tdk_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/generate-tdk"
                )
                with patch(
                    "services.server_tdk_export.LLMClient",
                    StubServerTdkLlm,
                ):
                    tdk_response = client.post(
                        tdk_path,
                        json={"revision": 2},
                    )
                self.assertEqual(
                    tdk_response.status_code,
                    200,
                    tdk_response.text,
                )
                tdk_delivered = tdk_response.json()
                self.assertEqual(tdk_delivered["revision"], 3)
                self.assertEqual(tdk_delivered["tdk_path"], "")
                self.assertTrue(tdk_delivered["tdk_asset_id"])
                self.assertEqual(
                    len(tdk_delivered["tdk_content_hash"]),
                    64,
                )
                self.assertEqual(
                    tdk_delivered["tdk_filename"],
                    "D.docx",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.images.prepared",
                        "article.docx.exported",
                        "article.tdk.generated",
                    ],
                )
                tdk_asset = (
                    PostgresKnowledgeAssetRepository(
                        self.engine
                    ).get_asset(
                        self.project_a,
                        tdk_delivered["tdk_asset_id"],
                    )
                )
                assert tdk_asset is not None
                tdk_key = str(tdk_asset.metadata["object_key"])
                tdk_document = Document(
                    BytesIO(private_store.objects[tdk_key])
                )
                self.assertEqual(
                    [paragraph.text[:3] for paragraph in tdk_document.paragraphs],
                    ["T: ", "D: ", "K: "],
                )
                self.assertEqual(
                    client.get(
                        (
                            f"/api/projects/{self.project_a}/assets/"
                            f"{tdk_delivered['tdk_asset_id']}/download"
                        )
                    ).status_code,
                    404,
                )
                tdk_download_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/tdk/download"
                )
                tdk_download = client.get(tdk_download_path)
                self.assertEqual(
                    tdk_download.status_code,
                    200,
                    tdk_download.text,
                )
                self.assertEqual(
                    tdk_download.json()["asset_id"],
                    tdk_delivered["tdk_asset_id"],
                )
                screenshot_bytes = self._image_bytes("white")
                screenshot_asset = (
                    object_service.upload_final_ai_screenshot(
                        actor=actor,
                        project_id=self.project_a,
                        asset_id=(
                            "asset_"
                            + hashlib.sha256(
                                screenshot_bytes
                            ).hexdigest()
                        ),
                        data=screenshot_bytes,
                        width=320,
                        height=240,
                    )
                )
                complete = repository.get(self.task_a)
                assert complete is not None
                complete["humanized_article"] = complete["linked_article"]
                complete["final_ai_check"] = {
                    "confirmed": True,
                    "screenshot_asset_id": screenshot_asset.asset_id,
                    "screenshot_content_hash": (
                        screenshot_asset.content_hash
                    ),
                    "screenshot_filename": "final-ai-rate.png",
                    "screenshot_width": 320,
                    "screenshot_height": 240,
                    "article_hash": hashlib.sha256(
                        complete["linked_article"].encode("utf-8")
                    ).hexdigest(),
                }
                repository.upsert(complete)
                package_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/package-delivery"
                )
                packaged = client.post(
                    package_path,
                    json={"revision": 3},
                )
                self.assertEqual(
                    packaged.status_code,
                    200,
                    packaged.text,
                )
                package_task = packaged.json()
                self.assertEqual(package_task["revision"], 4)
                self.assertEqual(
                    package_task["delivery_package_path"],
                    "",
                )
                self.assertTrue(
                    package_task["delivery_package_asset_id"]
                )
                self.assertEqual(
                    package_task["delivery_package_filename"],
                    (
                        f"{self.project_a}-topic_"
                        f"{package_task['topic_index']:03d}.zip"
                    ),
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.images.prepared",
                        "article.docx.exported",
                        "article.tdk.generated",
                        "article.delivery.packaged",
                    ],
                )
                package_asset = (
                    PostgresKnowledgeAssetRepository(
                        self.engine
                    ).get_asset(
                        self.project_a,
                        package_task[
                            "delivery_package_asset_id"
                        ],
                    )
                )
                assert package_asset is not None
                package_key = str(
                    package_asset.metadata["object_key"]
                )
                with ZipFile(
                    BytesIO(private_store.objects[package_key])
                ) as archive:
                    self.assertEqual(
                        set(archive.namelist()),
                        {
                            delivered["docx_filename"],
                            "D.docx",
                            *[
                                image["filename"]
                                for image in delivered["images"]
                            ],
                            "final-ai-rate.png",
                        },
                    )
                self.assertEqual(
                    client.get(
                        (
                            f"/api/projects/{self.project_a}/assets/"
                            f"{package_task['delivery_package_asset_id']}/"
                            "download"
                        )
                    ).status_code,
                    404,
                )
                package_download_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/delivery-package/download"
                )
                self.assertEqual(
                    client.get(package_download_path).status_code,
                    200,
                )
                put_count = len(private_store.put_calls)
                self.assertEqual(
                    client.post(
                        export_path,
                        json={"revision": 1},
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    len(private_store.put_calls),
                    put_count,
                )
                with patch(
                    "services.server_tdk_export.LLMClient",
                    StubServerTdkLlm,
                ):
                    self.assertEqual(
                        client.post(
                            tdk_path,
                            json={"revision": 2},
                        ).status_code,
                        409,
                    )
                self.assertEqual(
                    len(private_store.put_calls),
                    put_count,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="viewer")
                    )
                self.assertEqual(
                    client.get(docx_download_path).status_code,
                    403,
                )
                self.assertEqual(
                    client.get(tdk_download_path).status_code,
                    403,
                )
                self.assertEqual(
                    client.get(package_download_path).status_code,
                    403,
                )
                self.assertFalse(local_state.exists())

    def test_server_saves_reviewed_humanized_article_with_cas(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        initial = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "initial_ai_checked",
                "initial_article": initial,
                "article": initial,
                "humanized_article": "",
                "final_article": "stale final article",
                "article_versions": [],
            }
        )
        repository.upsert(record)
        candidate = initial.replace(
            "Keep the original application guidance.",
            "Use the reviewed application guidance.",
        )

        codec = ServerActorSessionCodec(b"h" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "h" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/humanized-article"
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "article": candidate,
                        },
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "article": candidate,
                            "status": "caller-controlled",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/humanized-article"
                        ),
                        json={
                            "revision": 0,
                            "article": candidate,
                        },
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "article": candidate.replace(
                                "## Buyer Checks",
                                "## Changed Buyer Checks",
                            ),
                        },
                    ).status_code,
                    422,
                )
                detector = RecordingAiRateDetector()
                with patch(
                    "server_project_http.ZeroGPTClient",
                    return_value=detector,
                ):
                    saved = client.put(
                        path,
                        json={
                            "revision": 0,
                            "article": candidate,
                            "recheck_ai_rate": True,
                        },
                    )
                self.assertEqual(saved.status_code, 200, saved.text)
                body = saved.json()
                self.assertEqual(body["revision"], 1)
                self.assertEqual(body["status"], "humanized_ready")
                self.assertEqual(body["humanized_article"], candidate)
                self.assertEqual(body["article"], candidate)
                self.assertEqual(detector.calls, [candidate])
                self.assertEqual(body["final_ai_check"]["provider"], "zerogpt")
                self.assertEqual(body["final_ai_check"]["score"], 18.5)
                self.assertEqual(
                    body["final_ai_check"]["article_hash"],
                    body["humanized_article_hash"],
                )
                self.assertEqual(body["final_article"], "")
                self.assertEqual(
                    body["article_versions"][-1]["kind"],
                    "humanized",
                )
                self.assertEqual(
                    body["article_versions"][-1]["source_kind"],
                    "external_manual",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.humanized.updated"],
                )
                self.assertNotIn(candidate, str(audit.events))

                class FailingDetector:
                    ready = True

                    def detect(self, text: str) -> ZeroGPTDetectionResult:
                        del text
                        raise RuntimeError("secret-provider-detail")

                revised_candidate = candidate.replace(
                    "Keep the original evidence guidance.",
                    "Use the reviewed evidence guidance.",
                )
                with patch(
                    "server_project_http.ZeroGPTClient",
                    return_value=FailingDetector(),
                ):
                    saved_without_detection = client.put(
                        path,
                        json={
                            "revision": 1,
                            "article": revised_candidate,
                            "recheck_ai_rate": True,
                        },
                    )
                self.assertEqual(
                    saved_without_detection.status_code,
                    200,
                    saved_without_detection.text,
                )
                fallback_body = saved_without_detection.json()
                self.assertEqual(fallback_body["revision"], 2)
                self.assertEqual(
                    fallback_body["humanized_article"],
                    revised_candidate,
                )
                self.assertIsNone(fallback_body["final_ai_check"]["score"])
                self.assertNotIn(
                    "secret-provider-detail",
                    fallback_body["final_ai_check"]["report"],
                )
                self.assertEqual(
                    client.put(
                        path,
                        json={
                            "revision": 0,
                            "article": candidate,
                        },
                    ).status_code,
                    409,
                )
                self.assertFalse(local_state.exists())

    def test_server_initial_ai_review_uses_private_screenshot_asset(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        initial = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "draft_ready",
                "raw_draft_article": initial,
                "initial_article": initial,
                "initial_article_hash": hashlib.sha256(
                    initial.encode("utf-8")
                ).hexdigest(),
                "article": initial,
                "initial_ai_check": {},
            }
        )
        repository.upsert(record)

        codec = ServerActorSessionCodec(b"i" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        private_store = FakeDownloadStore()
        object_service = ProjectKnowledgeObjectService(
            store=private_store,
            bucket="private-bucket",
            repository=PostgresKnowledgeAssetRepository(self.engine),
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "i" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.app.state.server_project_object_service = (
                    object_service
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                upload_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/initial-ai/screenshot"
                )
                confirm_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/initial-ai"
                )
                download_path = f"{upload_path}/download"
                self.assertEqual(
                    client.post(
                        upload_path,
                        params={"revision": 0},
                        files={
                            "file": (
                                "initial.png",
                                self._image_bytes("white"),
                                "image/png",
                            )
                        },
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="reviewer")
                    )
                self.assertEqual(
                    client.put(
                        confirm_path,
                        json={
                            "revision": 0,
                            "score": 12.5,
                            "report": "Reviewed initial AI result.",
                        },
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/checks/initial-ai/screenshot"
                        ),
                        params={"revision": 0},
                        files={
                            "file": (
                                "cross-project.png",
                                self._image_bytes("red"),
                                "image/png",
                            )
                        },
                    ).status_code,
                    403,
                )
                uploaded = client.post(
                    upload_path,
                    params={"revision": 0},
                    files={
                        "file": (
                            "initial.png",
                            self._image_bytes("white"),
                            "image/png",
                        )
                    },
                )
                self.assertEqual(
                    uploaded.status_code,
                    200,
                    uploaded.text,
                )
                screenshot_task = uploaded.json()
                self.assertEqual(screenshot_task["revision"], 1)
                self.assertEqual(
                    screenshot_task["status"],
                    "draft_ready",
                )
                check = screenshot_task["initial_ai_check"]
                self.assertEqual(check["screenshot_path"], "")
                self.assertTrue(check["screenshot_asset_id"])
                self.assertEqual(
                    check["screenshot_filename"],
                    "initial-ai-rate.png",
                )
                self.assertEqual(
                    (
                        check["screenshot_width"],
                        check["screenshot_height"],
                    ),
                    (320, 240),
                )
                download = client.get(download_path)
                self.assertEqual(download.status_code, 200, download.text)
                self.assertTrue(
                    download.json()["url"].startswith(
                        "https://signed.example.test/"
                    )
                )
                confirmed = client.put(
                    confirm_path,
                    json={
                        "revision": 1,
                        "score": 12.5,
                        "report": "Reviewed initial AI result.",
                        "confirmed": True,
                    },
                )
                self.assertEqual(
                    confirmed.status_code,
                    200,
                    confirmed.text,
                )
                confirmed_task = confirmed.json()
                self.assertEqual(confirmed_task["revision"], 2)
                self.assertEqual(
                    confirmed_task["status"],
                    "initial_ai_checked",
                )
                self.assertTrue(
                    confirmed_task["initial_ai_check"]["confirmed"]
                )
                self.assertEqual(
                    confirmed_task["initial_ai_check"]["article_hash"],
                    record["initial_article_hash"],
                )
                self.assertEqual(
                    confirmed_task["humanized_article"],
                    "",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.initial_ai_screenshot.uploaded",
                        "article.initial_ai_check.updated",
                    ],
                )
                self.assertNotIn(
                    "Reviewed initial AI result.",
                    str(audit.events),
                )
                self.assertEqual(
                    client.put(
                        confirm_path,
                        json={
                            "revision": 1,
                            "score": 12.5,
                            "report": "stale",
                        },
                    ).status_code,
                    409,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="viewer")
                    )
                self.assertEqual(
                    client.get(download_path).status_code,
                    403,
                )
                self.assertFalse(local_state.exists())

    def test_server_final_ai_review_uses_private_screenshot_asset(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "humanized_ready",
                "humanized_article": SERVER_ARTICLE,
                "article": SERVER_ARTICLE,
                "final_ai_check": {},
            }
        )
        repository.upsert(record)

        codec = ServerActorSessionCodec(b"v" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        private_store = FakeDownloadStore()
        object_service = ProjectKnowledgeObjectService(
            store=private_store,
            bucket="private-bucket",
            repository=PostgresKnowledgeAssetRepository(self.engine),
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "v" * 32,
                        "ARTICLE_AGENT_OBJECT_STORE_BUCKET": "",
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.app.state.server_project_object_service = (
                    object_service
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                upload_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/final-ai/screenshot"
                )
                confirm_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/checks/final-ai"
                )
                self.assertEqual(
                    client.post(
                        upload_path,
                        params={"revision": 0},
                        files={
                            "file": (
                                "final.png",
                                self._image_bytes("white"),
                                "image/png",
                            )
                        },
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="reviewer")
                    )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/checks/final-ai/screenshot"
                        ),
                        params={"revision": 0},
                        files={
                            "file": (
                                "cross-project.png",
                                self._image_bytes("red"),
                                "image/png",
                            )
                        },
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.put(
                        confirm_path,
                        json={
                            "revision": 0,
                            "score": 14.2,
                            "report": "Reviewed final AI result.",
                        },
                    ).status_code,
                    409,
                )

                uploaded = client.post(
                    upload_path,
                    params={"revision": 0},
                    files={
                        "file": (
                            "final.png",
                            self._image_bytes("white"),
                            "image/png",
                        )
                    },
                )

                self.assertEqual(
                    uploaded.status_code,
                    200,
                    uploaded.text,
                )
                screenshot_task = uploaded.json()
                self.assertEqual(screenshot_task["revision"], 1)
                self.assertEqual(
                    screenshot_task["status"],
                    "humanized_ready",
                )
                check = screenshot_task["final_ai_check"]
                self.assertEqual(check["screenshot_path"], "")
                self.assertTrue(check["screenshot_asset_id"])
                self.assertEqual(
                    len(check["screenshot_content_hash"]),
                    64,
                )
                self.assertEqual(
                    check["screenshot_filename"],
                    "final-ai-rate.png",
                )
                self.assertEqual(
                    (
                        check["screenshot_width"],
                        check["screenshot_height"],
                    ),
                    (None, None),
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.final_ai_screenshot.uploaded"],
                )
                self.assertNotIn(
                    "report",
                    audit.events[0].details,
                )
                replacement_bytes = b"replacement attachment without decoding"
                replaced = client.post(
                    upload_path,
                    params={"revision": 1},
                    files={
                        "file": (
                            "replacement.jpg",
                            replacement_bytes,
                            "image/jpeg",
                        )
                    },
                )
                self.assertEqual(replaced.status_code, 200, replaced.text)
                replaced_task = replaced.json()
                self.assertEqual(replaced_task["revision"], 2)
                self.assertEqual(
                    replaced_task["final_ai_check"]["screenshot_filename"],
                    "final-ai-rate.jpg",
                )
                self.assertIsNone(
                    replaced_task["final_ai_check"]["screenshot_width"]
                )
                self.assertIsNone(
                    replaced_task["final_ai_check"]["screenshot_height"]
                )
                self.assertFalse(replaced_task["final_ai_check"]["confirmed"])
                self.assertIn(replacement_bytes, private_store.objects.values())
                screenshot_task = replaced_task
                check = screenshot_task["final_ai_check"]
                asset_id = check["screenshot_asset_id"]
                self.assertEqual(
                    client.get(
                        (
                            f"/api/projects/{self.project_a}/assets/"
                            f"{asset_id}/download"
                        )
                    ).status_code,
                    404,
                )
                download_path = (
                    f"{upload_path}/download"
                )
                download = client.get(download_path)
                self.assertEqual(
                    download.status_code,
                    200,
                    download.text,
                )
                self.assertEqual(
                    download.json()["asset_id"],
                    asset_id,
                )

                confirmed = client.put(
                    confirm_path,
                    json={
                        "revision": 2,
                        "score": 14.2,
                        "report": "Reviewed final AI result.",
                    },
                )
                self.assertEqual(
                    confirmed.status_code,
                    200,
                    confirmed.text,
                )
                confirmed_task = confirmed.json()
                self.assertEqual(confirmed_task["revision"], 3)
                self.assertEqual(
                    confirmed_task["status"],
                    "final_ai_checked",
                )
                self.assertTrue(
                    confirmed_task["final_ai_check"]["confirmed"]
                )
                self.assertEqual(
                    confirmed_task["final_ai_check"]["report"],
                    "Reviewed final AI result.",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.final_ai_screenshot.uploaded",
                        "article.final_ai_screenshot.uploaded",
                        "article.final_ai_check.updated",
                    ],
                )
                self.assertNotIn(
                    "report",
                    audit.events[2].details,
                )
                linked_record = repository.get(self.task_a)
                assert linked_record is not None
                linked_record["status"] = "links_verified"
                linked_record["revision"] = 3
                linked_record["final_ai_check"]["confirmed"] = False
                repository.upsert(linked_record)
                reconfirmed = client.put(
                    confirm_path,
                    json={
                        "revision": 3,
                        "score": 13.8,
                        "report": "Rechecked after link restoration.",
                    },
                )
                self.assertEqual(reconfirmed.status_code, 200, reconfirmed.text)
                self.assertEqual(reconfirmed.json()["revision"], 4)
                self.assertEqual(
                    reconfirmed.json()["status"],
                    "links_verified",
                )
                self.assertTrue(
                    reconfirmed.json()["final_ai_check"]["confirmed"]
                )
                replaced_after_confirmation = client.post(
                    upload_path,
                    params={"revision": 4},
                    files={
                        "file": (
                            "newer.webp",
                            b"newer opaque bytes",
                            "image/webp",
                        )
                    },
                )
                self.assertEqual(
                    replaced_after_confirmation.status_code,
                    200,
                    replaced_after_confirmation.text,
                )
                replaced_after_confirmation_task = (
                    replaced_after_confirmation.json()
                )
                self.assertEqual(
                    replaced_after_confirmation_task["status"],
                    "links_verified",
                )
                self.assertFalse(
                    replaced_after_confirmation_task["final_ai_check"][
                        "confirmed"
                    ]
                )
                self.assertEqual(
                    replaced_after_confirmation_task["final_ai_check"][
                        "screenshot_filename"
                    ],
                    "final-ai-rate.webp",
                )
                put_count = len(private_store.put_calls)
                self.assertEqual(
                    client.post(
                        upload_path,
                        params={"revision": 0},
                        files={
                            "file": (
                                "stale.png",
                                self._image_bytes("black"),
                                "image/png",
                            )
                        },
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    len(private_store.put_calls),
                    put_count,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="viewer")
                    )
                self.assertEqual(
                    client.get(download_path).status_code,
                    403,
                )
                self.assertFalse(local_state.exists())

    def test_server_writing_settings_http_is_scoped_strict_and_safe(
        self,
    ) -> None:
        import app as app_module

        codec = ServerActorSessionCodec(b"w" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            audit = RecordingAuditWriter()
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "w" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                client.app.state.server_task_writing_settings_service_factory = (
                    ServerTaskWritingSettingsServiceFactory(
                        self.engine,
                        isolated,
                        audit=audit,
                    )
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                base_path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/writing-settings"
                )
                preview_payload = self._writing_settings_payload(
                    kind="outline"
                )

                self.assertEqual(
                    client.put(
                        f"/api/tasks/{self.task_a}/writing-settings",
                        json=self._writing_settings_payload(),
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.post(
                        f"/api/tasks/{self.task_a}/prompt-preview",
                        json={
                            "kind": "outline",
                            "selection": "system",
                        },
                    ).status_code,
                    404,
                )

                unknown = dict(preview_payload)
                unknown["prompt_content"] = "must not be accepted"
                self.assertEqual(
                    client.post(
                        f"{base_path}/preview",
                        json=unknown,
                    ).status_code,
                    422,
                )
                preview = client.post(
                    f"{base_path}/preview",
                    json=preview_payload,
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertEqual(
                    preview.headers.get("cache-control"),
                    "no-store",
                )
                body = preview.json()
                self.assertEqual(body["project_id"], self.project_a)
                self.assertEqual(body["task_id"], self.task_a)
                self.assertEqual(body["task_revision"], 0)
                self.assertEqual(body["kind"], "outline")
                self.assertEqual(
                    set(body["prompt_snapshot"]),
                    {
                        "prompt_id",
                        "name",
                        "kind",
                        "version",
                        "source",
                        "captured_at",
                    },
                )
                self.assertNotIn("content", body["prompt_snapshot"])
                self.assertNotIn("hash", str(body["prompt_snapshot"]))
                self.assertIsInstance(body["effective_prompt"], str)
                self.assertEqual(
                    body["warnings"],
                    [
                        "Preview resolves the current Project Prompt and "
                        "Published Knowledge; generation pins exact inputs "
                        "when the Job is enqueued."
                    ],
                )
                self.assertEqual(audit.events, [])

                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/writing-settings/preview"
                        ),
                        json=preview_payload,
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.put(
                        base_path,
                        json=self._writing_settings_payload(),
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )

                unknown_update = self._writing_settings_payload()
                unknown_update["prompt_content"] = "must not be accepted"
                self.assertEqual(
                    client.put(base_path, json=unknown_update).status_code,
                    422,
                )
                missing_update = self._writing_settings_payload()
                del missing_update["include_topic_notes"]
                self.assertEqual(
                    client.put(base_path, json=missing_update).status_code,
                    422,
                )
                non_boolean_update = self._writing_settings_payload()
                non_boolean_update["include_topic_notes"] = 1
                self.assertEqual(
                    client.put(base_path, json=non_boolean_update).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/writing-settings"
                        ),
                        json=self._writing_settings_payload(),
                    ).status_code,
                    403,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_a}/tasks/"
                            f"{self.task_b}/writing-settings"
                        ),
                        json=self._writing_settings_payload(),
                    ).status_code,
                    404,
                )

                invalid_selection = self._writing_settings_payload()
                invalid_selection["outline_prompt_selection"] = (
                    "private-secret-selection"
                )
                invalid = client.put(
                    base_path,
                    json=invalid_selection,
                )
                self.assertEqual(invalid.status_code, 422)
                self.assertNotIn("private-secret-selection", invalid.text)

                updated = client.put(
                    base_path,
                    json=self._writing_settings_payload(),
                )
                self.assertEqual(updated.status_code, 200, updated.text)
                self.assertEqual(updated.json()["revision"], 1)
                self.assertEqual(len(audit.events), 1)
                event = audit.events[0]
                self.assertEqual(
                    event.action,
                    "article.writing_settings.updated",
                )
                self.assertNotIn("Private topic notes", str(event.details))
                self.assertNotIn(
                    "Private outline instructions",
                    str(event.details),
                )
                self.assertNotIn(
                    "Private article instructions",
                    str(event.details),
                )
                self.assertEqual(
                    set(event.details),
                    {
                        "from_revision",
                        "to_revision",
                        "status",
                        "topic_notes_changed",
                        "outline_custom_prompt_changed",
                        "article_custom_prompt_changed",
                        "outline_prompt_selection_changed",
                        "article_prompt_selection_changed",
                        "use_outline_custom_prompt",
                        "use_article_custom_prompt",
                        "include_project_introduction",
                        "include_project_notes",
                        "include_topic_notes",
                        "outline_prompt_source",
                        "outline_prompt_version",
                        "article_prompt_source",
                        "article_prompt_version",
                    },
                )
                self.assertEqual(
                    client.put(
                        base_path,
                        json=self._writing_settings_payload(),
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_a}/tasks/"
                            "missing/writing-settings/preview"
                        ),
                        json=preview_payload,
                    ).status_code,
                    404,
                )
                self.assertEqual(
                    client.post(
                        f"{base_path}/preview",
                        json=preview_payload,
                    ).status_code,
                    409,
                )
                failed_payload = self._writing_settings_payload(revision=1)
                failed_payload["topic_notes"] = "Must roll back"
                client.app.state.server_task_writing_settings_service_factory = (
                    ServerTaskWritingSettingsServiceFactory(
                        self.engine,
                        isolated,
                        audit=FailingAuditWriter(),
                    )
                )
                failed = client.put(base_path, json=failed_payload)
                self.assertEqual(failed.status_code, 503, failed.text)
                self.assertNotIn("private audit failure", failed.text)
                stored = client.get(
                    f"/api/projects/{self.project_a}/tasks/{self.task_a}"
                )
                self.assertEqual(stored.status_code, 200, stored.text)
                self.assertEqual(stored.json()["revision"], 1)
                self.assertEqual(
                    stored.json()["topic_notes"],
                    "Private topic notes",
                )
                self.assertFalse(local_state.exists())

    def test_server_title_generation_uses_system_template_and_published_scope(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "new",
                "selected_title": "",
                "title_candidates": [],
                "outline": "stale confirmed outline",
                "outline_draft": "stale outline draft",
                "article": "stale article",
            }
        )
        repository.upsert(record)
        published_chunk = self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-title-published",
            text="Topic 2 has one verified published positioning fact.",
        )
        self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-title-blog",
            text="Topic 2 blog positioning must not influence titles.",
            source_kind="official_blog",
            trust_tier="reference_material",
        )
        self._store_outline_context(
            project_id=self.project_b,
            suffix=f"{self.task_b}-title-published",
            text="Topic 2 repeated is a cross-project positioning fact.",
        )
        provider = RecordingTitleProvider()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        handler = ServerTitleGenerationHandler(
            self.engine,
            provider=provider,
            audit=audit,
        )
        codec = ServerActorSessionCodec(b"t" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            registry = ServerTitleGenerationRegistry(
                self.engine,
                config=isolated,
                access=access,
                handler=handler,
                audit=audit,
            )
            self.addCleanup(registry.stop)
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "t" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_title_generation = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/titles"
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "instruction": "client override",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/titles"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(public_job["operation"], "titles")
                self.assertNotIn("request", public_job)
                self.assertNotIn("requested_by_user_id", public_job)
                status_path = (
                    f"{path}/jobs/{public_job['job_id']}"
                )
                terminal = None
                for _attempt in range(100):
                    response = client.get(status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(terminal["status"], "succeeded")
                self.assertEqual(terminal["result_revision"], 1)
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["chunk_ids"],
                    [published_chunk],
                )
                stored_payload = repository.get(self.task_a)
                assert stored_payload is not None
                stored = TaskRecord.model_validate(stored_payload)
                self.assertEqual(stored.revision, 1)
                self.assertEqual(stored.status, "titles_ready")
                self.assertEqual(len(stored.title_candidates), 10)
                self.assertEqual(stored.selected_title, "")
                self.assertEqual(stored.outline, "")
                self.assertEqual(stored.outline_draft, "")
                self.assertEqual(stored.article, "")
                self.assertNotIn(
                    "Server candidate title",
                    str(audit.events),
                )
                self.assertFalse(local_state.exists())

    def test_server_outline_generation_uses_pinned_prompt_and_published_scope(
        self,
    ) -> None:
        import app as app_module

        published_chunk = self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-published",
            text=(
                "Selected topic 2 uses one verified published project fact."
            ),
        )
        self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-inbox",
            text=(
                "Selected topic 2 must not use this unpublished project fact."
            ),
            status="inbox",
        )
        self._store_outline_context(
            project_id=self.project_b,
            suffix=f"{self.task_b}-published",
            text=(
                "Selected topic 2 repeated repeated is a cross-project fact."
            ),
        )
        provider = RecordingOutlineProvider()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        handler = ServerOutlineGenerationHandler(
            self.engine,
            provider=provider,
            audit=audit,
        )
        registry = ServerOutlineGenerationRegistry(
            self.engine,
            access=access,
            handler=handler,
            audit=audit,
        )
        self.addCleanup(registry.stop)
        codec = ServerActorSessionCodec(b"o" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "o" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_outline_generation = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/outline"
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "prompt_id": "client-controlled",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/outline"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(public_job["operation"], "outline")
                self.assertNotIn("request", public_job)
                self.assertNotIn("requested_by_user_id", public_job)
                status_path = (
                    f"{path}/jobs/{public_job['job_id']}"
                )
                terminal = None
                for _attempt in range(100):
                    response = client.get(status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(terminal["status"], "succeeded")
                self.assertEqual(terminal["result_revision"], 1)
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["chunk_ids"],
                    [published_chunk],
                )
                self.assertEqual(
                    provider.calls[0]["prompt_source"],
                    "system",
                )
                stored_payload = self._task_repository(
                    organization_id=self.org_a,
                    project_id=self.project_a,
                ).get(self.task_a)
                assert stored_payload is not None
                stored = TaskRecord.model_validate(stored_payload)
                self.assertEqual(stored.revision, 1)
                self.assertEqual(stored.outline, "")
                self.assertEqual(
                    stored.outline_draft,
                    "## Generated server outline\n\n"
                    "### Evidence-led section",
                )
                self.assertEqual(
                    stored.article_versions[-1].kind,
                    "outline_draft",
                )
                self.assertEqual(
                    stored.article_versions[-1].source_kind,
                    "generated",
                )
                assert stored.last_outline_prompt_snapshot is not None
                self.assertEqual(
                    stored.last_outline_prompt_snapshot.source,
                    "system",
                )
                self.assertNotIn(
                    "Generated server outline",
                    str(audit.events),
                )
                self.assertFalse(local_state.exists())

    def test_server_article_generation_uses_pinned_prompt_and_published_scope(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "outline_confirmed",
                "selected_title": "Example Buyer Guide",
                "outline": (
                    "## Buyer Checks\n\n"
                    "### Confirm the application\n\n"
                    "### Compare evidence\n\n"
                    "## FAQ"
                ),
                "outline_draft": (
                    "## Buyer Checks\n\n"
                    "### Confirm the application\n\n"
                    "### Compare evidence\n\n"
                    "## FAQ"
                ),
                "raw_draft_article": "",
                "initial_article": "",
                "article": "",
                "humanized_article": "stale downstream article",
                "article_versions": [],
            }
        )
        repository.upsert(record)
        published_chunk = self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-article-published",
            text=(
                "Example Buyer Guide for Topic 2 uses one verified "
                "published project fact."
            ),
        )
        blog_chunk = self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-article-blog",
            text=(
                "Example Buyer Guide for Topic 2 includes one relevant "
                "editorial comparison."
            ),
            source_kind="official_blog",
            trust_tier="reference_material",
        )
        self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-article-inbox",
            text=(
                "Example Buyer Guide for Topic 2 must not use this "
                "unpublished fact."
            ),
            status="inbox",
        )
        self._store_outline_context(
            project_id=self.project_b,
            suffix=f"{self.task_b}-article-published",
            text=(
                "Example Buyer Guide for Topic 2 repeated repeated "
                "is a cross-project fact."
            ),
        )
        provider = RecordingArticleProvider()
        ai_rate = RecordingAiRateDetector()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        handler = ServerArticleGenerationHandler(
            self.engine,
            provider=provider,
            ai_rate=ai_rate,
            audit=audit,
        )
        codec = ServerActorSessionCodec(b"a" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            registry = ServerArticleGenerationRegistry(
                self.engine,
                config=isolated,
                access=access,
                handler=handler,
                audit=audit,
            )
            self.addCleanup(registry.stop)
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "a" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_article_generation = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article"
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "word_count": 1000,
                        },
                    ).status_code,
                    422,
                )
                rewrite_path = f"{path}/rewrite"
                self.assertEqual(
                    client.post(
                        rewrite_path,
                        json={"revision": 0},
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/article"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(public_job["operation"], "article")
                self.assertNotIn("request", public_job)
                self.assertNotIn("requested_by_user_id", public_job)
                status_path = f"{path}/jobs/{public_job['job_id']}"
                terminal = None
                for _attempt in range(100):
                    response = client.get(status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(terminal["status"], "succeeded")
                self.assertEqual(terminal["result_revision"], 1)
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["chunk_ids"],
                    [published_chunk, blog_chunk],
                )
                self.assertEqual(
                    provider.calls[0]["prompt_source"],
                    "system",
                )
                self.assertEqual(
                    provider.calls[0]["target_words"],
                    1200,
                )
                stored_payload = repository.get(self.task_a)
                assert stored_payload is not None
                stored = TaskRecord.model_validate(stored_payload)
                self.assertEqual(stored.revision, 1)
                self.assertEqual(stored.status, "draft_ready")
                self.assertEqual(
                    stored.raw_draft_article,
                    SERVER_ARTICLE.strip(),
                )
                self.assertEqual(
                    stored.initial_article,
                    SERVER_ARTICLE.strip(),
                )
                self.assertEqual(len(ai_rate.calls), 1)
                self.assertEqual(stored.initial_ai_check.provider, "zerogpt")
                self.assertEqual(stored.initial_ai_check.score, 18.5)
                self.assertEqual(
                    stored.initial_ai_check.article_hash,
                    content_hash(stored.initial_article),
                )
                self.assertFalse(stored.initial_ai_check.confirmed)
                self.assertEqual(stored.article, stored.initial_article)
                self.assertEqual(stored.humanized_article, "")
                self.assertEqual(
                    [item.kind for item in stored.article_versions],
                    ["raw_draft", "initial"],
                )
                assert stored.last_article_prompt_snapshot is not None
                self.assertEqual(
                    stored.last_article_prompt_snapshot.source,
                    "system",
                )
                self.assertNotIn(
                    "Example Buyer Guide",
                    str(audit.events),
                )
                self.assertFalse(local_state.exists())
                self.assertEqual(
                    client.post(path, json={"revision": 1}).status_code,
                    409,
                )
                with (
                    patch(
                        "services.server_article_generation._latest_completed_research_context",
                        side_effect=AssertionError(
                            "disabled Evidence Pack lookup should not run"
                        ),
                    ),
                    patch.object(
                        PostgresPublishedOutlineContext,
                        "select",
                        side_effect=AssertionError(
                            "disabled Evidence Pack generation must not search"
                        ),
                    ),
                ):
                    rewritten = client.post(
                        rewrite_path,
                        json={
                            "revision": 1,
                            "use_evidence_pack": False,
                        },
                    )
                self.assertEqual(
                    rewritten.status_code,
                    200,
                    rewritten.text,
                )
                rewrite_job = rewritten.json()
                self.assertEqual(
                    rewrite_job["operation"],
                    "rewrite_article",
                )
                self.assertNotIn("request", rewrite_job)
                rewrite_status_path = (
                    f"{rewrite_path}/jobs/{rewrite_job['job_id']}"
                )
                rewrite_terminal = None
                for _attempt in range(100):
                    response = client.get(rewrite_status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    rewrite_terminal = response.json()
                    if rewrite_terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert rewrite_terminal is not None
                self.assertEqual(
                    rewrite_terminal["status"],
                    "succeeded",
                )
                self.assertEqual(rewrite_terminal["result_revision"], 2)
                with self.engine.connect() as connection:
                    rewrite_request = connection.execute(
                        sa.select(background_jobs.c.request).where(
                            background_jobs.c.organization_id == self.org_a,
                            background_jobs.c.project_id == self.project_a,
                            background_jobs.c.job_id
                            == rewrite_job["job_id"],
                        )
                    ).scalar_one()
                self.assertFalse(rewrite_request["use_evidence_pack"])
                rewritten_payload = repository.get(self.task_a)
                assert rewritten_payload is not None
                rewritten_task = TaskRecord.model_validate(
                    rewritten_payload
                )
                self.assertEqual(rewritten_task.revision, 2)
                self.assertEqual(len(provider.calls), 2)
                self.assertEqual(len(ai_rate.calls), 2)
                self.assertEqual(rewritten_task.initial_ai_check.score, 18.5)
                self.assertFalse(rewritten_task.initial_ai_check.confirmed)
                self.assertEqual(provider.calls[1]["chunk_ids"], [])
                self.assertEqual(
                    [item.kind for item in rewritten_task.article_versions],
                    ["raw_draft", "initial", "raw_draft", "initial"],
                )
                self.assertEqual(
                    rewritten_task.article_versions[-1].source_kind,
                    "regenerated_raw_draft",
                )
                self.assertIn(
                    "article.article_regeneration.queued",
                    [event.action for event in audit.events],
                )
                self.assertIn(
                    "article.draft.regenerated",
                    [event.action for event in audit.events],
                )

    def test_outline_context_rejects_unpublished_pinned_chunk(
        self,
    ) -> None:
        suffix = f"{self.task_a}-context-revoked"
        chunk_id = self._store_outline_context(
            project_id=self.project_a,
            suffix=suffix,
            text="Selected topic 2 published context.",
        )
        context = PostgresPublishedOutlineContext(self.engine)
        selected = context.select(
            project_id=self.project_a,
            query="Selected topic 2",
        )
        self.assertEqual(
            [chunk.chunk_id for chunk in selected],
            [chunk_id],
        )
        with self.engine.begin() as connection:
            connection.execute(
                knowledge_sources.update()
                .where(
                    knowledge_sources.c.project_id == self.project_a,
                    knowledge_sources.c.source_id == f"{suffix}-source",
                )
                .values(
                    status="stale",
                    current_snapshot_id=None,
                )
            )
        with self.assertRaisesRegex(
            JobConflict,
            "^published outline context changed$",
        ):
            context.load_current(
                project_id=self.project_a,
                chunk_ids=(chunk_id,),
            )

    def test_server_link_restoration_is_project_scoped_and_hash_bound(
        self,
    ) -> None:
        import app as app_module

        repository, source, candidate = self._prepare_link_task()
        source_hash = content_hash(source)
        candidate_hash = content_hash(candidate)

        provider = RecordingLinkRestorationProvider()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        handler = ServerLinkRestorationHandler(
            self.engine,
            provider=provider,
            audit=audit,
        )
        registry = ServerLinkRestorationRegistry(
            self.engine,
            access=access,
            handler=handler,
            audit=audit,
        )
        self.addCleanup(registry.stop)
        codec = ServerActorSessionCodec(b"l" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "l" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_link_restoration = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/restore-links"
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "article": "caller-controlled",
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/restore-links"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(
                    public_job["operation"],
                    "restore_links",
                )
                self.assertNotIn("request", public_job)
                self.assertNotIn("requested_by_user_id", public_job)
                status_path = f"{path}/jobs/{public_job['job_id']}"
                terminal = None
                for _attempt in range(100):
                    response = client.get(status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(
                    terminal["status"],
                    "succeeded",
                    terminal,
                )
                self.assertEqual(terminal["result_revision"], 1)
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["source_hash"],
                    source_hash,
                )
                self.assertEqual(
                    provider.calls[0]["candidate_hash"],
                    candidate_hash,
                )
                stored_payload = repository.get(self.task_a)
                assert stored_payload is not None
                stored = TaskRecord.model_validate(stored_payload)
                self.assertEqual(stored.status, "links_verified")
                self.assertEqual(stored.article, source)
                self.assertEqual(stored.linked_article, source)
                self.assertTrue(stored.link_validation.passed)
                self.assertEqual(
                    stored.article_versions[-1].kind,
                    "linked",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.link_restoration.queued",
                        "article.links.restored",
                        "background_job.terminal",
                    ],
                )
                self.assertNotIn(source, str(audit.events))
                self.assertNotIn(candidate, str(audit.events))
                self.assertFalse(local_state.exists())

    def test_server_saves_seo_review_settings_with_prompt_validation(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        codec = ServerActorSessionCodec(b"v" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "v" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-review-settings"
                )
                request = {
                    "revision": 0,
                    "primary_keyword": "  buyer   guide ",
                    "long_tail_keywords": [
                        "fastener   sourcing",
                        "FASTENER SOURCING",
                        "quality checks",
                    ],
                    "prompt_selection": "system",
                }
                self.assertEqual(
                    client.put(path, json=request).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                escaped = dict(request)
                escaped["prompt_content"] = "caller prompt"
                self.assertEqual(
                    client.put(path, json=escaped).status_code,
                    422,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/seo-review-settings"
                        ),
                        json=request,
                    ).status_code,
                    403,
                )
                response = client.put(path, json=request)
                self.assertEqual(response.status_code, 200, response.text)
                body = response.json()
                self.assertEqual(body["revision"], 1)
                self.assertEqual(
                    body["seo_primary_keyword"],
                    "buyer guide",
                )
                self.assertEqual(
                    body["seo_long_tail_keywords"],
                    ["fastener sourcing", "quality checks"],
                )
                self.assertEqual(
                    body["seo_review_prompt_selection"],
                    "system",
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.seo_review_settings.updated"],
                )
                self.assertEqual(
                    audit.events[0].details[
                        "long_tail_keyword_count"
                    ],
                    2,
                )
                self.assertNotIn("buyer guide", str(audit.events))
                self.assertNotIn("fastener sourcing", str(audit.events))
                self.assertEqual(
                    client.put(path, json=request).status_code,
                    409,
                )
                stored = repository.get(self.task_a)
                assert stored is not None
                self.assertEqual(stored["revision"], 1)
                self.assertFalse(local_state.exists())

    def test_server_seo_review_uses_pinned_prompt_and_published_scope(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "draft_ready",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "seo_primary_keyword": "buyer guide",
                "seo_long_tail_keywords": ["supplier checks"],
                "seo_review_prompt_selection": "system",
                "seo_reviews": [],
            }
        )
        repository.upsert(record)
        published_chunk = self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-review-published",
            text=(
                "Selected topic 2 Topic 2 buyer guide supplier checks "
                "published SEO review evidence."
            ),
        )
        self._store_outline_context(
            project_id=self.project_a,
            suffix=f"{self.task_a}-review-inbox",
            text="Unpublished review evidence.",
            status="inbox",
        )
        self._store_outline_context(
            project_id=self.project_b,
            suffix=f"{self.task_b}-review-published",
            text="Cross-project review evidence.",
        )
        provider = RecordingSeoReviewProvider()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        handler = ServerSeoReviewGenerationHandler(
            self.engine,
            provider=provider,
            audit=audit,
        )
        registry = ServerSeoReviewGenerationRegistry(
            self.engine,
            access=access,
            handler=handler,
            audit=audit,
        )
        self.addCleanup(registry.stop)
        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_seo_review_generation = (
                    registry
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/seo-reviews"
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="reviewer")
                    )
                self.assertEqual(
                    client.post(
                        path,
                        json={
                            "revision": 0,
                            "prompt_snapshot": {"content": "attacker"},
                        },
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/seo-reviews"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(public_job["operation"], "seo_review")
                self.assertNotIn("request", public_job)
                self.assertNotIn("requested_by_user_id", public_job)
                status_path = f"{path}/jobs/{public_job['job_id']}"
                terminal = None
                for _attempt in range(100):
                    response = client.get(status_path)
                    self.assertEqual(
                        response.status_code,
                        200,
                        response.text,
                    )
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(
                    terminal["status"],
                    "succeeded",
                    terminal,
                )
                self.assertEqual(terminal["result_revision"], 1)
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["chunk_ids"],
                    [published_chunk],
                )
                self.assertNotIn(
                    "Unpublished review evidence.",
                    str(provider.calls),
                )
                self.assertNotIn(
                    "Cross-project review evidence.",
                    str(provider.calls),
                )
                stored_payload = repository.get(self.task_a)
                assert stored_payload is not None
                stored = TaskRecord.model_validate(stored_payload)
                self.assertEqual(stored.revision, 1)
                self.assertEqual(stored.status, "draft_ready")
                self.assertEqual(stored.initial_article, article)
                self.assertEqual(len(stored.seo_reviews), 1)
                self.assertEqual(stored.seo_reviews[0].status, "open")
                self.assertTrue(stored.seo_reviews[0].publish_ready)
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.seo_review.queued",
                        "article.seo_review.generated",
                        "background_job.terminal",
                    ],
                )
                self.assertNotIn(article, str(audit.events))
                self.assertNotIn("No blocking issue.", str(audit.events))
                self.assertFalse(local_state.exists())

    def test_server_humanize_uses_pinned_project_prompt(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "initial_ai_checked",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
            }
        )
        repository.upsert(record)
        actor = ActorIdentity(self.org_a, self.user_a)
        prompt = PromptSnapshot(
            prompt_id="server-humanize-v1",
            name="Server Humanize",
            kind="humanize",
            content="Rewrite safely.\n\n{{ARTICLE}}",
            version=1,
            source="project_default",
            captured_at="2026-07-31T00:00:00+00:00",
        )
        provider = RecordingHumanizeProvider()
        detector = RecordingAiRateDetector()
        audit = RecordingAuditWriter()
        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        registry = ServerHumanizeGenerationRegistry(
            self.engine,
            access=access,
            handler=ServerHumanizeGenerationHandler(
                self.engine,
                provider=provider,
                ai_rate=detector,
                audit=audit,
            ),
            audit=audit,
        )
        self.addCleanup(registry.stop)
        codec = ServerActorSessionCodec(b"z" * 32)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.object(
                    PostgresProjectPromptService,
                    "resolve",
                    return_value=prompt,
                ),
                patch(
                    "services.server_humanize_generation."
                    "load_pinned_project_prompt",
                    return_value=prompt,
                ),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_humanize_generation = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/humanize"
                )
                self.assertEqual(
                    client.post(
                        path,
                        json={"revision": 0, "prompt_id": "attacker"},
                    ).status_code,
                    422,
                )
                self.assertEqual(
                    client.post(path, json={"revision": 0}).status_code,
                    403,
                )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/humanize"
                        ),
                        json={"revision": 0},
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                queued = client.post(path, json={"revision": 0})
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertNotIn("request", public_job)
                terminal = None
                for _attempt in range(100):
                    response = client.get(
                        f"{path}/jobs/{public_job['job_id']}"
                    )
                    self.assertEqual(response.status_code, 200)
                    terminal = response.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(
                    terminal["status"],
                    "succeeded",
                    terminal,
                )
                self.assertEqual(len(provider.calls), 1)
                self.assertEqual(
                    provider.calls[0]["prompt_id"],
                    prompt.prompt_id,
                )
                stored = repository.get(self.task_a)
                assert stored is not None
                self.assertEqual(stored["revision"], 1)
                self.assertEqual(stored["status"], "humanized_ready")
                self.assertEqual(stored["humanized_article"], article)
                self.assertEqual(detector.calls, [article])
                self.assertEqual(
                    stored["final_ai_check"]["provider"],
                    "zerogpt",
                )
                self.assertEqual(
                    stored["final_ai_check"]["score"],
                    18.5,
                )
                self.assertEqual(
                    stored["final_ai_check"]["article_hash"],
                    stored["humanized_article_hash"],
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.humanize.queued",
                        "article.humanized.generated",
                        "background_job.terminal",
                    ],
                )
                self.assertNotIn(article, str(audit.events))
                self.assertFalse(local_state.exists())

    def test_server_humanize_requires_explicit_project_default(
        self,
    ) -> None:
        _repository, _article = self._prepare_humanize_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        registry = ServerHumanizeGenerationRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=None,
        )
        self.addCleanup(registry.stop)

        with self.assertRaisesRegex(
            JobConflict,
            "project humanize prompt is not configured",
        ):
            registry.enqueue(
                actor=ActorIdentity(self.org_a, self.user_a),
                project_id=self.project_a,
                task_id=self.task_a,
                source_revision=0,
            )

        with self.engine.connect() as connection:
            queued = connection.execute(
                sa.select(sa.func.count())
                .select_from(background_jobs)
                .where(
                    background_jobs.c.organization_id == self.org_a,
                    background_jobs.c.project_id == self.project_a,
                    background_jobs.c.operation == "humanize",
                )
            ).scalar_one()
        self.assertEqual(queued, 0)

    def test_humanize_worker_rejects_source_article_drift(
        self,
    ) -> None:
        repository, article = self._prepare_humanize_task()
        prompt = PromptSnapshot(
            prompt_id="server-humanize-v1",
            name="Server Humanize",
            kind="humanize",
            content="Rewrite safely.\n\n{{ARTICLE}}",
            version=1,
            source="project_default",
            captured_at="2026-07-31T00:00:00+00:00",
        )
        reference = ProjectPromptReference.from_snapshot(prompt)
        job = {
            "operation": "humanize",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "task_id": self.task_a,
            "requested_by_user_id": self.user_a,
            "source_revision": 0,
            "request": {
                **reference.private_values(),
                "source_article_hash": content_hash(article),
            },
        }
        changed = repository.get(self.task_a)
        assert changed is not None
        changed_article = article.replace(
            "Keep the original evidence guidance.",
            "Keep updated evidence guidance.",
        )
        changed["initial_article"] = changed_article
        changed["initial_article_hash"] = content_hash(changed_article)
        changed["article"] = changed_article
        repository.upsert(changed)
        provider = RecordingHumanizeProvider()
        handler = ServerHumanizeGenerationHandler(
            self.engine,
            provider=provider,
        )

        with self.assertRaisesRegex(
            JobConflict,
            "source article changed",
        ):
            handler(job, lambda: False)

        self.assertEqual(provider.calls, [])
        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(persisted["humanized_article"], "")

    def test_humanize_prompt_resolution_errors_are_sanitized(
        self,
    ) -> None:
        self._prepare_humanize_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        registry = ServerHumanizeGenerationRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=None,
        )
        self.addCleanup(registry.stop)
        secret = "private-prompt-store-detail"

        with (
            patch.object(
                PostgresProjectPromptService,
                "resolve",
                side_effect=RuntimeError(secret),
            ),
            self.assertRaises(
                HumanizeGenerationUnavailable
            ) as raised,
        ):
            registry.enqueue(
                actor=ActorIdentity(self.org_a, self.user_a),
                project_id=self.project_a,
                task_id=self.task_a,
                source_revision=0,
            )

        self.assertNotIn(secret, str(raised.exception))

    def test_humanize_audit_failure_rolls_back_task(self) -> None:
        repository, article = self._prepare_humanize_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        prompt = PromptSnapshot(
            prompt_id="server-humanize-v1",
            name="Server Humanize",
            kind="humanize",
            content="Rewrite safely.\n\n{{ARTICLE}}",
            version=1,
            source="project_default",
            captured_at="2026-07-31T00:00:00+00:00",
        )
        reference = ProjectPromptReference.from_snapshot(prompt)
        handler = ServerHumanizeGenerationHandler(
            self.engine,
            provider=RecordingHumanizeProvider(),
            audit=FailingAuditWriter(),
        )
        job = {
            "operation": "humanize",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "task_id": self.task_a,
            "requested_by_user_id": self.user_a,
            "source_revision": 0,
            "request": {
                **reference.private_values(),
                "source_article_hash": content_hash(article),
            },
        }

        with (
            patch(
                "services.server_humanize_generation."
                "load_pinned_project_prompt",
                return_value=prompt,
            ),
            self.assertRaises(ServerTaskCommandUnavailable) as raised,
        ):
            handler(job, lambda: False)

        self.assertNotIn("private audit failure", str(raised.exception))
        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(persisted["status"], "initial_ai_checked")
        self.assertEqual(persisted["humanized_article"], "")
        self.assertEqual(persisted["article_versions"], [])

    def test_humanize_worker_reauthorizes_before_provider_call(
        self,
    ) -> None:
        _repository, article = self._prepare_humanize_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        prompt = PromptSnapshot(
            prompt_id="server-humanize-v1",
            name="Server Humanize",
            kind="humanize",
            content="Rewrite safely.\n\n{{ARTICLE}}",
            version=1,
            source="project_default",
            captured_at="2026-07-31T00:00:00+00:00",
        )
        reference = ProjectPromptReference.from_snapshot(prompt)
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        batch = raw_queue.create_batch(
            "humanize",
            [
                {
                    "task_id": self.task_a,
                    "source_revision": 0,
                    "customer": self.project_a,
                    "topic_index": 2,
                    "request": {
                        **reference.private_values(),
                        "source_article_hash": content_hash(article),
                    },
                }
            ],
            customer=self.project_a,
            requested_by_user_id=self.user_a,
        )
        job_id = str(batch["jobs"][0]["id"])
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="viewer")
            )
        provider = RecordingHumanizeProvider()
        registry = ServerHumanizeGenerationRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=ServerHumanizeGenerationHandler(
                self.engine,
                provider=provider,
            ),
            audit=RecordingAuditWriter(),
        )
        self.addCleanup(registry.stop)

        registry.start_existing()

        current = None
        for _attempt in range(100):
            current = raw_queue.get_job(job_id)
            if current["status"] == "conflict":
                break
            time.sleep(0.02)
        assert current is not None
        self.assertEqual(current["status"], "conflict")
        self.assertEqual(
            current["error"],
            "job actor is not authorized",
        )
        self.assertEqual(provider.calls, [])

    def test_seo_review_worker_reauthorizes_before_provider_call(
        self,
    ) -> None:
        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "draft_ready",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "seo_review_prompt_selection": "system",
                "seo_reviews": [],
            }
        )
        repository.upsert(record)
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="reviewer")
            )
        prompt = ProjectPromptReference.from_snapshot(
            PromptSnapshot(
                kind="review",
                source="system",
                captured_at="2026-07-31T00:00:00+00:00",
            )
        )
        template = ReviewTemplateReference.current()
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        batch = raw_queue.create_batch(
            "seo_review",
            [
                {
                    "task_id": self.task_a,
                    "source_revision": 0,
                    "customer": self.project_a,
                    "topic_index": 2,
                    "request": {
                        **prompt.private_values(),
                        **template.private_values(),
                        "context_chunk_ids": [],
                        "source_article_hash": content_hash(article),
                    },
                }
            ],
            customer=self.project_a,
            requested_by_user_id=self.user_a,
        )
        job_id = str(batch["jobs"][0]["id"])
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="viewer")
            )
        provider = RecordingSeoReviewProvider()
        registry = ServerSeoReviewGenerationRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=ServerSeoReviewGenerationHandler(
                self.engine,
                provider=provider,
            ),
            audit=RecordingAuditWriter(),
        )
        self.addCleanup(registry.stop)

        registry.start_existing()

        current = None
        for _attempt in range(100):
            current = raw_queue.get_job(job_id)
            if current["status"] == "conflict":
                break
            time.sleep(0.02)
        assert current is not None
        self.assertEqual(current["status"], "conflict")
        self.assertEqual(
            current["error"],
            "job actor is not authorized",
        )
        self.assertEqual(provider.calls, [])

    def test_seo_review_audit_failure_rolls_back_review_run(
        self,
    ) -> None:
        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        record.update(
            {
                "status": "draft_ready",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "seo_review_prompt_selection": "system",
                "seo_reviews": [],
            }
        )
        repository.upsert(record)
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="reviewer")
            )

        class FailingAudit:
            def append(self, connection, event) -> None:
                del connection, event
                raise RuntimeError("private audit failure")

        prompt = ProjectPromptReference.from_snapshot(
            PromptSnapshot(
                kind="review",
                source="system",
                captured_at="2026-07-31T00:00:00+00:00",
            )
        )
        template = ReviewTemplateReference.current()
        handler = ServerSeoReviewGenerationHandler(
            self.engine,
            provider=RecordingSeoReviewProvider(),
            audit=FailingAudit(),
        )
        job = {
            "id": "job-seo-audit-failure",
            "operation": "seo_review",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "task_id": self.task_a,
            "requested_by_user_id": self.user_a,
            "source_revision": 0,
            "request": {
                **prompt.private_values(),
                **template.private_values(),
                "context_chunk_ids": [],
                "source_article_hash": content_hash(article),
            },
        }

        with self.assertRaises(ServerTaskCommandUnavailable) as raised:
            handler(job, lambda: False)

        self.assertNotIn(
            "private audit failure",
            str(raised.exception),
        )
        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(persisted["status"], "draft_ready")
        self.assertEqual(persisted["seo_reviews"], [])

    def test_server_seo_review_decision_preview_and_apply_are_scoped(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        target = "Keep the original evidence guidance."
        proposed = "Compare the original evidence before approval."
        start = article.index(target)
        review = SeoReviewRun(
            id="review-server-a",
            source_article=article,
            source_article_hash=content_hash(article),
            source_revision=0,
            score=80,
            dimensions=[
                SeoReviewDimension(
                    key="intent",
                    name="Intent",
                    score=8,
                    target_score=9,
                )
            ],
            report="Private review report.",
            changes=[
                SeoReviewChange(
                    id="change-server-a",
                    operation="replace",
                    title="Clarify evidence",
                    target_text=target,
                    model_proposed_text=proposed,
                    reviewed_text=proposed,
                    source_start=start,
                    source_end=start + len(target),
                )
            ],
            prompt_snapshot=PromptSnapshot(
                kind="review",
                source="system",
                content="Private rubric.",
            ),
            created_at="2026-07-31T00:00:00+00:00",
        )
        record.update(
            {
                "status": "draft_ready",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "seo_reviews": [review.model_dump(mode="json")],
            }
        )
        repository.upsert(record)
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="reviewer")
            )
        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    app_module.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                base = (
                    f"/api/projects/{self.project_a}/tasks/{self.task_a}"
                    "/seo-reviews/review-server-a"
                )
                invalid = client.put(
                    f"{base}/changes/change-server-a",
                    json={
                        "revision": 0,
                        "decision": "accepted",
                        "reviewed_text": proposed,
                        "review_id": "attacker",
                    },
                )
                self.assertEqual(invalid.status_code, 422)
                decided = client.put(
                    f"{base}/changes/change-server-a",
                    json={
                        "revision": 0,
                        "decision": "accepted",
                        "reviewed_text": proposed,
                    },
                )
                self.assertEqual(decided.status_code, 200, decided.text)
                self.assertEqual(decided.json()["revision"], 1)
                preview = client.post(
                    f"{base}/preview",
                    json={"revision": 1},
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertIn(proposed, preview.json()["article"])
                apply_payload = {
                    "revision": 1,
                    "preview_hash": preview.json()["article_hash"],
                    "confirm_pending": True,
                }
                self.assertEqual(
                    client.post(
                        f"{base}/apply",
                        json=apply_payload,
                    ).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                applied = client.post(
                    f"{base}/apply",
                    json=apply_payload,
                )
                self.assertEqual(applied.status_code, 200, applied.text)
                body = applied.json()
                self.assertEqual(body["revision"], 2)
                self.assertEqual(
                    body["seo_reviews"][0]["status"],
                    "applied",
                )
                self.assertIn(proposed, body["initial_article"])
                self.assertEqual(body["status"], "draft_ready")
                self.assertEqual(
                    [event.action for event in audit.events],
                    [
                        "article.seo_review.change.updated",
                        "article.seo_review.applied",
                    ],
                )
                self.assertNotIn(
                    "Private review report.",
                    str(audit.events),
                )
                self.assertNotIn(proposed, str(audit.events))
                self.assertEqual(
                    client.post(
                        f"{base}/complete",
                        json={
                            "revision": 2,
                            "confirm_pending": True,
                        },
                    ).status_code,
                    409,
                )
    def test_server_seo_review_complete_rolls_back_on_audit_failure(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        article = SERVER_ARTICLE.strip()
        review = SeoReviewRun(
            id="review-server-rollback",
            source_article=article,
            source_article_hash=content_hash(article),
            source_revision=0,
            score=90,
            report="Private report.",
            changes=[],
            prompt_snapshot=PromptSnapshot(
                kind="review",
                source="system",
                content="Private rubric.",
            ),
            created_at="2026-07-31T00:00:00+00:00",
        )
        record.update(
            {
                "status": "draft_ready",
                "initial_article": article,
                "initial_article_hash": content_hash(article),
                "article": article,
                "seo_reviews": [review.model_dump(mode="json")],
            }
        )
        repository.upsert(record)
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="reviewer")
            )
        codec = ServerActorSessionCodec(b"z" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "z" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_project_task_store_factory = (
                    ServerProjectTaskStoreFactory(
                        self.engine,
                        isolated,
                        audit=FailingAuditWriter(),
                    )
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                response = client.post(
                    (
                        f"/api/projects/{self.project_a}/tasks/"
                        f"{self.task_a}/seo-reviews/"
                        "review-server-rollback/complete"
                    ),
                    json={"revision": 0, "confirm_pending": True},
                )
                self.assertEqual(response.status_code, 503, response.text)
                self.assertNotIn("private audit failure", response.text)

        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(
            persisted["seo_reviews"][0]["status"],  # type: ignore[index]
            "open",
        )

    def test_link_worker_rejects_article_hash_drift_before_provider(
        self,
    ) -> None:
        repository, source, candidate = self._prepare_link_task()
        template = LinkTemplateReference.current()
        job = {
            "operation": "restore_links",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "task_id": self.task_a,
            "requested_by_user_id": self.user_a,
            "source_revision": 0,
            "request": {
                **template.private_values(),
                "source_article_hash": content_hash(source),
                "candidate_article_hash": content_hash(candidate),
                "source_link_count": 1,
            },
        }
        changed = repository.get(self.task_a)
        assert changed is not None
        changed_candidate = candidate.replace(
            "Keep the original evidence guidance.",
            "Keep revised evidence guidance.",
        )
        changed["humanized_article"] = changed_candidate
        changed["humanized_article_hash"] = content_hash(
            changed_candidate
        )
        changed["final_ai_check"]["article_hash"] = content_hash(  # type: ignore[index]
            changed_candidate
        )
        repository.upsert(changed)
        provider = RecordingLinkRestorationProvider()
        handler = ServerLinkRestorationHandler(
            self.engine,
            provider=provider,
        )

        with self.assertRaisesRegex(
            JobConflict,
            "candidate article changed",
        ):
            handler(job, lambda: False)

        self.assertEqual(provider.calls, [])
        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(persisted["status"], "final_ai_checked")

    def test_link_restoration_audit_failure_rolls_back_task(self) -> None:
        repository, source, candidate = self._prepare_link_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )

        class FailingAudit:
            def append(self, connection, event) -> None:
                del connection, event
                raise RuntimeError("audit unavailable")

        template = LinkTemplateReference.current()
        handler = ServerLinkRestorationHandler(
            self.engine,
            provider=RecordingLinkRestorationProvider(),
            audit=FailingAudit(),
        )
        job = {
            "operation": "restore_links",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "task_id": self.task_a,
            "requested_by_user_id": self.user_a,
            "source_revision": 0,
            "request": {
                **template.private_values(),
                "source_article_hash": content_hash(source),
                "candidate_article_hash": content_hash(candidate),
                "source_link_count": 1,
            },
        }

        with self.assertRaises(ServerTaskCommandUnavailable):
            handler(job, lambda: False)

        persisted = repository.get(self.task_a)
        assert persisted is not None
        self.assertEqual(persisted["revision"], 0)
        self.assertEqual(persisted["status"], "final_ai_checked")
        self.assertEqual(persisted["linked_article"], "")

    def test_link_worker_reauthorizes_before_provider_call(self) -> None:
        _, source, candidate = self._prepare_link_task()
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        template = LinkTemplateReference.current()
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        batch = raw_queue.create_batch(
            "restore_links",
            [
                {
                    "task_id": self.task_a,
                    "source_revision": 0,
                    "customer": self.project_a,
                    "topic_index": 2,
                    "request": {
                        **template.private_values(),
                        "source_article_hash": content_hash(source),
                        "candidate_article_hash": content_hash(candidate),
                        "source_link_count": 1,
                    },
                }
            ],
            customer=self.project_a,
            requested_by_user_id=self.user_a,
        )
        job_id = str(batch["jobs"][0]["id"])
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="viewer")
            )
        provider = RecordingLinkRestorationProvider()
        registry = ServerLinkRestorationRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=ServerLinkRestorationHandler(
                self.engine,
                provider=provider,
            ),
            audit=RecordingAuditWriter(),
        )
        self.addCleanup(registry.stop)

        registry.start_existing()

        current = None
        for _attempt in range(100):
            current = raw_queue.get_job(job_id)
            if current["status"] == "conflict":
                break
            time.sleep(0.02)
        assert current is not None
        self.assertEqual(current["status"], "conflict")
        self.assertEqual(
            current["error"],
            "job actor is not authorized",
        )
        self.assertEqual(provider.calls, [])

    def test_server_product_rediscovery_uses_requested_actor_and_pg_worker(
        self,
    ) -> None:
        import app as app_module

        calls: list[dict[str, object]] = []

        def handler(job, cancelled):
            self.assertFalse(cancelled())
            calls.append(
                {
                    "organization_id": job["organization_id"],
                    "project_id": job["project_id"],
                    "task_id": job["task_id"],
                    "requested_by_user_id": job[
                        "requested_by_user_id"
                    ],
                    "request": dict(job["request"]),
                }
            )
            return int(job["source_revision"])

        access = ProjectAccessService(
            PostgresProjectAccessRepository(self.engine)
        )
        rediscovery_audit = RecordingAuditWriter()
        registry = ServerProductRediscoveryRegistry(
            self.engine,
            access=access,
            handler=handler,
            audit=rediscovery_audit,
        )
        self.addCleanup(registry.stop)
        codec = ServerActorSessionCodec(b"r" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        task_repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        task_before = task_repository.get(self.task_a)
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "r" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                app_module.app.state.server_product_rediscovery = registry
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/product-rediscovery"
                )
                payload = {
                    "revision": 0,
                    "category_url": (
                        f"https://{self.project_a}/products"
                    ),
                    "max_products": 3,
                }
                self.assertEqual(
                    client.post(path, json=payload).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                self.assertEqual(
                    client.post(
                        (
                            f"/api/projects/{self.project_b}/tasks/"
                            f"{self.task_b}/product-rediscovery"
                        ),
                        json=payload,
                    ).status_code,
                    403,
                )
                queued = client.post(path, json=payload)
                self.assertEqual(queued.status_code, 200, queued.text)
                public_job = queued.json()
                self.assertEqual(
                    public_job["operation"],
                    "product_rediscovery",
                )
                self.assertNotIn("request", public_job)
                self.assertNotIn(
                    "requested_by_user_id",
                    public_job,
                )
                self.assertNotIn("error", public_job)

                status_path = (
                    f"{path}/jobs/{public_job['job_id']}"
                )
                terminal = None
                for _attempt in range(100):
                    current = client.get(status_path)
                    self.assertEqual(
                        current.status_code,
                        200,
                        current.text,
                    )
                    terminal = current.json()
                    if terminal["status"] in {
                        "succeeded",
                        "failed",
                        "conflict",
                        "cancelled",
                    }:
                        break
                    time.sleep(0.02)
                assert terminal is not None
                self.assertEqual(terminal["status"], "succeeded")
                self.assertEqual(terminal["result_revision"], 0)
                self.assertEqual(
                    client.get(
                        (
                            f"/api/projects/{self.project_a}/tasks/"
                            f"{self.task_b}/product-rediscovery/jobs/"
                            f"{public_job['job_id']}"
                        )
                    ).status_code,
                    404,
                )
                self.assertEqual(len(calls), 1)
                self.assertEqual(
                    calls[0],
                    {
                        "organization_id": self.org_a,
                        "project_id": self.project_a,
                        "task_id": self.task_a,
                        "requested_by_user_id": self.user_a,
                        "request": {
                            "category_url": (
                                f"https://{self.project_a}/products"
                            ),
                            "max_products": 3,
                        },
                    },
                )
                with self.engine.connect() as connection:
                    requested_by = connection.execute(
                        sa.select(
                            background_jobs.c.requested_by_user_id
                        ).where(
                            background_jobs.c.organization_id
                            == self.org_a,
                            background_jobs.c.project_id
                            == self.project_a,
                            background_jobs.c.job_id
                            == public_job["job_id"],
                        )
                    ).scalar_one()
                self.assertEqual(requested_by, self.user_a)
                self.assertEqual(len(rediscovery_audit.events), 2)
                audit_event = rediscovery_audit.events[0]
                self.assertEqual(
                    audit_event.action,
                    "knowledge.products.rediscovery.queued",
                )
                self.assertEqual(
                    audit_event.target_id,
                    public_job["job_id"],
                )
                self.assertEqual(
                    audit_event.details,
                    {
                        "operation": "product_rediscovery",
                        "source_revision": 0,
                        "max_products": 3,
                    },
                )
                self.assertNotIn(
                    payload["category_url"],
                    str(audit_event.details),
                )
                terminal_audit = rediscovery_audit.events[1]
                self.assertEqual(
                    terminal_audit.action,
                    "background_job.terminal",
                )
                self.assertEqual(
                    terminal_audit.target_id,
                    public_job["job_id"],
                )
                self.assertEqual(
                    terminal_audit.details,
                    {
                        "operation": "product_rediscovery",
                        "status": "succeeded",
                        "attempts": 1,
                        "source_revision": 0,
                        "result_revision": 0,
                    },
                )
                self.assertNotIn(
                    payload["category_url"],
                    str(terminal_audit.details),
                )
                self.assertEqual(
                    client.post(
                        path,
                        json={**payload, "revision": 99},
                    ).status_code,
                    409,
                )
                self.assertEqual(
                    task_repository.get(self.task_a),
                    task_before,
                )
                self.assertFalse(local_state.exists())
            registry.stop()

    def test_product_rediscovery_audit_failure_rolls_back_job(self) -> None:
        class FailingAudit:
            def append(self, connection, event):
                raise RuntimeError(
                    "audit provider included https://secret.example.test"
                )

        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        registry = ServerProductRediscoveryRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=lambda job, cancelled: 0,
            audit=FailingAudit(),
        )
        self.addCleanup(registry.stop)
        command = ProductRediscoveryCommand(
            category_url="https://secret.example.test/products",
            max_products=3,
        )

        with self.assertRaisesRegex(
            ProductRediscoveryUnavailable,
            "^product rediscovery could not be queued$",
        ) as caught:
            registry.enqueue(
                actor=ActorIdentity(self.org_a, self.user_a),
                project_id=self.project_a,
                task_id=self.task_a,
                source_revision=0,
                command=command,
            )

        self.assertNotIn("secret.example.test", str(caught.exception))
        with self.engine.connect() as connection:
            job_count = connection.execute(
                sa.select(sa.func.count())
                .select_from(background_jobs)
                .where(
                    background_jobs.c.organization_id == self.org_a,
                    background_jobs.c.project_id == self.project_a,
                    background_jobs.c.task_id == self.task_a,
                )
            ).scalar_one()
            batch_count = connection.execute(
                sa.select(sa.func.count())
                .select_from(job_batches)
                .where(
                    job_batches.c.organization_id == self.org_a,
                    job_batches.c.project_id == self.project_a,
                )
            ).scalar_one()
        self.assertEqual(job_count, 0)
        self.assertEqual(batch_count, 0)

    def test_product_rediscovery_stop_requeues_and_reports_drain(self) -> None:
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        entered = threading.Event()

        def cooperative_handler(job, cancelled):
            entered.set()
            while not cancelled():
                time.sleep(0.005)
            raise JobCancelled("server is stopping")

        registry = ServerProductRediscoveryRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=cooperative_handler,
            audit=RecordingAuditWriter(),
        )
        job = registry.enqueue(
            actor=ActorIdentity(self.org_a, self.user_a),
            project_id=self.project_a,
            task_id=self.task_a,
            source_revision=0,
            command=ProductRediscoveryCommand(
                category_url=f"https://{self.project_a}/products",
                max_products=3,
            ),
        )
        self.assertTrue(entered.wait(timeout=2))

        report = registry.stop(timeout_seconds=2)

        self.assertTrue(report.drained)
        self.assertEqual(report.project_runner_count, 1)
        self.assertEqual(report.remaining_jobs, 0)
        self.assertEqual(registry.stop(), report)
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        interrupted = raw_queue.get_job(str(job["job_id"]))
        self.assertEqual(interrupted["status"], "queued")
        self.assertFalse(interrupted["cancel_requested"])

        replacement = ServerProductRediscoveryRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=lambda queued, cancelled: int(
                queued["source_revision"]
            ),
            audit=RecordingAuditWriter(),
        )
        self.addCleanup(replacement.stop)
        replacement.start_existing()
        deadline = time.time() + 2
        while time.time() < deadline:
            resumed = raw_queue.get_job(str(job["job_id"]))
            if resumed["status"] == "succeeded":
                break
            time.sleep(0.01)
        else:
            self.fail("Interrupted product rediscovery did not resume.")

    def test_product_rediscovery_handler_uses_active_project_domain(
        self,
    ) -> None:
        calls: list[dict[str, object]] = []

        class FakeSync:
            def sync_category(self, **kwargs):
                calls.append(dict(kwargs))
                return object()

        handler = ServerProductRediscoveryHandler(
            self.engine,
            sync_factory=lambda actor, project_id, job_id, cancelled: (
                calls.append(
                    {
                        "factory_organization_id": actor.organization_id,
                        "factory_user_id": actor.user_id,
                        "factory_project_id": project_id,
                        "factory_job_id": job_id,
                        "factory_cancelled": cancelled(),
                    }
                )
                or FakeSync()
            ),
        )
        job = {
            "operation": "product_rediscovery",
            "organization_id": self.org_a,
            "project_id": self.project_a,
            "id": "job-product-rediscovery",
            "requested_by_user_id": self.user_a,
            "task_id": self.task_a,
            "source_revision": 0,
            "request": {
                "category_url": (
                    f"https://{self.project_a}/products"
                ),
                "max_products": 4,
            },
        }

        result = handler(job, lambda: False)

        self.assertEqual(result, 0)
        self.assertEqual(
            calls,
            [
                {
                    "factory_organization_id": self.org_a,
                    "factory_user_id": self.user_a,
                    "factory_project_id": self.project_a,
                    "factory_job_id": "job-product-rediscovery",
                    "factory_cancelled": False,
                },
                {
                    "project_id": self.project_a,
                    "site_url": f"https://{self.project_a}",
                    "category_url": (
                        f"https://{self.project_a}/products"
                    ),
                    "max_products": 4,
                },
            ],
        )
        stale = dict(job)
        stale["source_revision"] = 99
        with self.assertRaisesRegex(
            JobConflict,
            "source task revision changed",
        ):
            handler(stale, lambda: False)
        self.assertEqual(len(calls), 2)

    def test_product_rediscovery_status_survives_unconfigured_runner(
        self,
    ) -> None:
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        batch = raw_queue.create_batch(
            "product_rediscovery",
            [
                {
                    "task_id": self.task_a,
                    "source_revision": 0,
                    "customer": self.project_a,
                    "topic_index": 2,
                    "request": {
                        "category_url": (
                            f"https://{self.project_a}/products"
                        ),
                        "max_products": 3,
                    },
                }
            ],
            customer=self.project_a,
            requested_by_user_id=self.user_a,
        )
        job_id = str(batch["jobs"][0]["id"])
        registry = ServerProductRediscoveryRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=None,
        )
        self.addCleanup(registry.stop)
        actor = ActorIdentity(self.org_a, self.user_a)

        current = registry.get_job(
            actor=actor,
            project_id=self.project_a,
            task_id=self.task_a,
            job_id=job_id,
        )

        self.assertEqual(current["status"], "queued")
        self.assertIsNone(current["started_at"])
        with self.assertRaisesRegex(
            ProductRediscoveryUnavailable,
            "not configured",
        ):
            registry.enqueue(
                actor=actor,
                project_id=self.project_a,
                task_id=self.task_a,
                source_revision=0,
                command=ProductRediscoveryCommand(
                    category_url=(
                        f"https://{self.project_a}/products"
                    ),
                    max_products=3,
                ),
            )
        with self.engine.begin() as connection:
            connection.execute(
                background_jobs.update()
                .where(
                    background_jobs.c.organization_id == self.org_a,
                    background_jobs.c.project_id == self.project_a,
                    background_jobs.c.job_id == job_id,
                )
                .values(operation="outline")
            )
        with self.assertRaises(KeyError):
            registry.get_job(
                actor=actor,
                project_id=self.project_a,
                task_id=self.task_a,
                job_id=job_id,
            )

    def test_product_rediscovery_worker_rejects_revoked_requester(
        self,
    ) -> None:
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="editor")
            )
        raw_queue = PostgresJobQueue(
            self.engine,
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        batch = raw_queue.create_batch(
            "product_rediscovery",
            [
                {
                    "task_id": self.task_a,
                    "source_revision": 0,
                    "customer": self.project_a,
                    "topic_index": 2,
                    "request": {
                        "category_url": (
                            f"https://{self.project_a}/products"
                        ),
                        "max_products": 3,
                    },
                }
            ],
            customer=self.project_a,
            requested_by_user_id=self.user_a,
        )
        job_id = str(batch["jobs"][0]["id"])
        with self.engine.begin() as connection:
            connection.execute(
                project_memberships.update()
                .where(
                    project_memberships.c.organization_id == self.org_a,
                    project_memberships.c.project_id == self.project_a,
                    project_memberships.c.user_id == self.user_a,
                )
                .values(role="viewer")
            )
        calls: list[str] = []
        registry = ServerProductRediscoveryRegistry(
            self.engine,
            access=ProjectAccessService(
                PostgresProjectAccessRepository(self.engine)
            ),
            handler=lambda job, cancelled: (
                calls.append(str(job["id"])) or 0
            ),
            audit=RecordingAuditWriter(),
        )
        self.addCleanup(registry.stop)

        registry.start_existing()

        current = None
        for _attempt in range(100):
            current = raw_queue.get_job(job_id)
            if current["status"] == "conflict":
                break
            time.sleep(0.02)
        assert current is not None
        self.assertEqual(current["status"], "conflict")
        self.assertEqual(
            current["error"],
            "job actor is not authorized",
        )
        self.assertEqual(calls, [])

    def test_server_rewrites_only_one_snapshotted_article_section(
        self,
    ) -> None:
        import app as app_module

        repository = self._task_repository(
            organization_id=self.org_a,
            project_id=self.project_a,
        )
        record = repository.get(self.task_a)
        assert record is not None
        record.update(
            {
                "status": "humanized_ready",
                "selected_title": "Example Buyer Guide",
                "initial_article": SERVER_ARTICLE,
                "humanized_article": "downstream copy",
                "article": "downstream copy",
                "article_versions": [],
            }
        )
        repository.upsert(record)

        codec = ServerActorSessionCodec(b"q" * 32)
        actor = ActorIdentity(self.org_a, self.user_a)
        base_config = app_module.config()
        with tempfile.TemporaryDirectory() as directory:
            local_state = Path(directory) / "must-not-exist"
            isolated = replace(
                base_config,
                knowledge_agent_enabled=False,
            )
            with (
                patch.object(app_module, "config", return_value=isolated),
                patch.dict(
                    os.environ,
                    {
                        "ARTICLE_AGENT_SERVER_MODE": "true",
                        "ARTICLE_AGENT_SERVER_SESSION_SECRET": "q" * 32,
                    },
                    clear=False,
                ),
                TestClient(app_module.app) as client,
            ):
                audit = self._install_recording_audit(
                    client.app,
                    isolated,
                )
                client.cookies.set(
                    SERVER_AUTH_COOKIE_NAME,
                    codec.create(actor),
                )
                path = (
                    f"/api/projects/{self.project_a}/tasks/"
                    f"{self.task_a}/article/sections"
                )
                request = {
                    "revision": 0,
                    "heading_path": ["Buyer Checks"],
                    "replacement_body": (
                        "### Confirm the application\n\n"
                        "Use revised application guidance.\n\n"
                        "### Compare evidence\n\n"
                        "Use revised evidence guidance."
                    ),
                }
                self.assertEqual(
                    client.put(path, json=request).status_code,
                    403,
                )
                with self.engine.begin() as connection:
                    connection.execute(
                        project_memberships.update()
                        .where(
                            project_memberships.c.organization_id
                            == self.org_a,
                            project_memberships.c.project_id
                            == self.project_a,
                            project_memberships.c.user_id == self.user_a,
                        )
                        .values(role="editor")
                    )
                escaped = dict(request)
                escaped["replacement_body"] = (
                    "## FAQ\n\nAttempted sibling replacement."
                )
                self.assertEqual(
                    client.put(path, json=escaped).status_code,
                    422,
                )
                self.assertEqual(
                    repository.get(self.task_a)["revision"],  # type: ignore[index]
                    0,
                )
                self.assertEqual(
                    client.put(
                        (
                            f"/api/projects/{self.project_a}/tasks/"
                            f"{self.task_b}/article/sections"
                        ),
                        json=request,
                    ).status_code,
                    404,
                )

                response = client.put(path, json=request)
                self.assertEqual(response.status_code, 200, response.text)
                saved = response.json()
                self.assertEqual(saved["revision"], 1)
                self.assertEqual(saved["status"], "draft_ready")
                self.assertEqual(saved["humanized_article"], "")
                self.assertIn(
                    "Use revised application guidance.",
                    saved["initial_article"],
                )
                self.assertNotIn(
                    "Keep the original application guidance.",
                    saved["initial_article"],
                )
                original_prefix = SERVER_ARTICLE.split(
                    "## Buyer Checks",
                    1,
                )[0]
                original_suffix = (
                    "## FAQ" + SERVER_ARTICLE.split("## FAQ", 1)[1]
                )
                self.assertTrue(
                    saved["initial_article"].startswith(original_prefix)
                )
                self.assertTrue(
                    saved["initial_article"].endswith(original_suffix)
                )
                self.assertEqual(
                    [
                        item["source_kind"]
                        for item in saved["article_versions"]
                    ],
                    [
                        "before_section_rewrite",
                        "section_rewrite",
                    ],
                )
                self.assertEqual(
                    saved["article_versions"][0]["content"],
                    SERVER_ARTICLE,
                )
                self.assertEqual(
                    [event.action for event in audit.events],
                    ["article.section.replaced"],
                )
                self.assertEqual(
                    client.put(path, json=request).status_code,
                    409,
                )
                persisted = repository.get(self.task_a)
                self.assertEqual(
                    len(persisted["article_versions"]),  # type: ignore[index]
                    2,
                )
                self.assertFalse(local_state.exists())


if __name__ == "__main__":
    unittest.main()
